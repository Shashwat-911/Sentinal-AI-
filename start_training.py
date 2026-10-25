import os
import sys
import time
import shutil
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from ultralytics import YOLO

def init_excel_log(save_dir, run_name):
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Training Log"

        headers = ["Epoch", "Box Loss", "mAP@50", "Timestamp"]
        header_fill = PatternFill("solid", fgColor="0A0A0A")
        header_font = Font(name="Consolas", bold=True, color="00FF88", size=11)

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col)].width = 18

        ws_meta = wb.create_sheet("Run Info")
        ws_meta["A1"] = "SENTINEL AI — Training Run"
        ws_meta["A2"] = f"Run Name: {run_name}"
        ws_meta["A3"] = f"Started: {time.strftime('%Y-%m-%d %H:%M:%S')}"
        ws_meta["A4"] = "Model: yolov9s"
        ws_meta["A5"] = "Epochs: 30"
        ws_meta["A6"] = "Image Size: 640"
        ws_meta["A7"] = "Batch Size: 4"

        os.makedirs(save_dir, exist_ok=True)
        excel_path = os.path.join(save_dir, f"{run_name}_training_log.xlsx")
        wb.save(excel_path)
        print(f"[EXCEL] Log started: {excel_path}")
        return excel_path
    except Exception as e:
        print(f"[EXCEL] Error creating log: {e}")
        return None

def append_excel_row(excel_path, epoch, loss, mAP):
    if not excel_path or not os.path.isfile(excel_path):
        return
    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["Training Log"]
        row = ws.max_row + 1
        fill_color = "1A1A1A" if row % 2 == 0 else "111111"
        fill = PatternFill("solid", fgColor=fill_color)
        font = Font(name="Consolas", color="E0E0E0", size=10)

        values = [epoch, round(loss, 6), round(mAP, 6), time.strftime('%H:%M:%S')]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.fill = fill
            cell.font = font
        wb.save(excel_path)
    except Exception as e:
        print(f"[EXCEL] Row append failed: {e}")

def export_graphs(loss_epochs, train_losses, map_epochs, map_values, graphs_dir):
    try:
        os.makedirs(graphs_dir, exist_ok=True)
        ts = time.strftime("%Y%m%d_%H%M%S")

        if loss_epochs and train_losses:
            fig, ax = plt.subplots(figsize=(10, 5), facecolor="#0a0a0a")
            ax.set_facecolor("#1a1a1a")
            ax.plot(loss_epochs, train_losses, color="#00ff88", linewidth=2, marker="o", markersize=4)
            ax.fill_between(loss_epochs, train_losses, alpha=0.1, color="#00ff88")
            ax.set_title("SENTINEL AI — Training Loss", color="#00ff88", fontsize=14, fontfamily="monospace")
            ax.set_xlabel("Epoch", color="#666666", fontfamily="monospace")
            ax.set_ylabel("Box Loss", color="#666666", fontfamily="monospace")
            ax.tick_params(colors="#666666")
            for spine in ax.spines.values():
                spine.set_color("#666666")
            ax.grid(True, alpha=0.15, color="#666666")
            loss_path = os.path.join(graphs_dir, f"loss_curve_{ts}.png")
            fig.savefig(loss_path, dpi=150, bbox_inches="tight", facecolor="#0a0a0a")
            plt.close(fig)
            print(f"[GRAPH] Loss curve saved: {loss_path}")

        if map_epochs and map_values:
            fig, ax = plt.subplots(figsize=(10, 5), facecolor="#0a0a0a")
            ax.set_facecolor("#1a1a1a")
            ax.plot(map_epochs, map_values, color="#ffaa00", linewidth=2, marker="o", markersize=4)
            ax.fill_between(map_epochs, map_values, alpha=0.1, color="#ffaa00")
            ax.set_title("SENTINEL AI — mAP@50", color="#ffaa00", fontsize=14, fontfamily="monospace")
            ax.set_xlabel("Epoch", color="#666666", fontfamily="monospace")
            ax.set_ylabel("mAP@50", color="#666666", fontfamily="monospace")
            ax.tick_params(colors="#666666")
            for spine in ax.spines.values():
                spine.set_color("#666666")
            ax.grid(True, alpha=0.15, color="#666666")
            map_path = os.path.join(graphs_dir, f"map_curve_{ts}.png")
            fig.savefig(map_path, dpi=150, bbox_inches="tight", facecolor="#0a0a0a")
            plt.close(fig)
            print(f"[GRAPH] mAP curve saved: {map_path}")
    except Exception as e:
        print(f"[GRAPH] Export failed: {e}")

def main():
    print("=" * 60)
    print("SENTINEL AI — STARTING 30-EPOCH TRAINING RUN")
    print("=" * 60)
    data_yaml = os.path.join(BASE_DIR, "data.yaml")
    save_dir = os.path.join(BASE_DIR, "models")
    run_name = "train"

    excel_path = init_excel_log(save_dir, run_name)

    model = YOLO("yolov9s.pt")
    
    print("\nStarting model.train()...\n")
    results = model.train(
        data=data_yaml,
        epochs=30,
        imgsz=640,
        batch=4,
        workers=0,
        project=save_dir,
        name=run_name,
        exist_ok=True,
        verbose=True
    )

    print("\nTraining completed!")
    
    # Extract mAP@50 and metrics
    metrics = results.results_dict
    map50 = metrics.get("metrics/mAP50(B)", 0.0)
    print(f"\n============================================================")
    print(f"FINAL TRAINING RESULTS")
    print(f"============================================================")
    print(f"  Final mAP@50: {map50:.5f}")

    # Ensure the best checkpoint is copied to the conventional model output paths.
    best_weights = os.path.join(save_dir, run_name, "weights", "best.pt")
    target_best = os.path.join(save_dir, "best.pt")
    if os.path.isfile(best_weights):
        shutil.copy2(best_weights, target_best)
        print(f"  Saved best weights to: {target_best}")

    # Read CSV results if available to populate Excel log and graphs
    results_csv = os.path.join(save_dir, run_name, "results.csv")
    if os.path.isfile(results_csv):
        import csv
        loss_epochs, train_losses, map_epochs, map_values = [], [], [], []
        with open(results_csv, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                clean_row = {k.strip(): v.strip() for k, v in row.items()}
                epoch_val = int(clean_row.get("epoch", 0)) + 1
                box_loss = float(clean_row.get("train/box_loss", 0.0))
                map_val = float(clean_row.get("metrics/mAP50(B)", 0.0))
                
                loss_epochs.append(epoch_val)
                train_losses.append(box_loss)
                map_epochs.append(epoch_val)
                map_values.append(map_val)
                
                append_excel_row(excel_path, epoch_val, box_loss, map_val)

        export_graphs(loss_epochs, train_losses, map_epochs, map_values, save_dir)

    print("============================================================")

if __name__ == "__main__":
    main()

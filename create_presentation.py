import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # LIGHT PROFESSIONAL ACADEMIC THEME
    # =========================================================================
    COLOR_BG          = RGBColor(248, 250, 252) # #F8FAFC (Soft Crisp Light Background)
    COLOR_CARD        = RGBColor(255, 255, 255) # #FFFFFF (Pure White Card Containers)
    COLOR_BORDER      = RGBColor(226, 232, 240) # #E2E8F0 (Subtle Slate Border)
    COLOR_HEADER      = RGBColor(15, 23, 42)    # #0F172A (Deep Navy / Slate Main Text)
    COLOR_BODY        = RGBColor(51, 65, 85)    # #334155 (Dark Slate Body Text)
    COLOR_MUTED       = RGBColor(100, 116, 139) # #64748B (Muted Gray Text)
    
    # Accent Colors
    COLOR_TEAL        = RGBColor(13, 148, 136)  # #0D9488 (Emerald Teal Accent)
    COLOR_BLUE        = RGBColor(37, 99, 235)   # #2563EB (Vibrant Royal Blue)
    COLOR_AMBER       = RGBColor(217, 119, 6)   # #D97706 (Warm Amber)
    COLOR_PURPLE      = RGBColor(124, 58, 237)  # #7C3AED (Deep Violet)

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    def add_header(slide, title_text, category_text="SENTINEL AI — PUBLIC SAFETY SURVEILLANCE SYSTEM"):
        # Header category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_TEAL
        p_cat.font.name = "Consolas"

        # Main Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.75))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_HEADER
        p_title.font.name = "Arial"

        # Subtle divider line under header
        div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.02))
        div.fill.solid()
        div.fill.fore_color.rgb = COLOR_BORDER
        div.line.color.rgb = COLOR_BORDER

    # =========================================================================
    # SLIDE 1: Title Slide (Light Theme)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Left decorative vertical accent bar
    line = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.0), Inches(0.15), Inches(3.6))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_TEAL
    line.line.color.rgb = COLOR_TEAL

    # Title & Subtitle Container Card
    card_title = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.9), Inches(11.3), Inches(3.8))
    card_title.fill.solid()
    card_title.fill.fore_color.rgb = COLOR_CARD
    card_title.line.color.rgb = COLOR_BORDER

    tf1 = card_title.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.4)
    tf1.margin_top = Inches(0.4)

    p0 = tf1.paragraphs[0]
    p0.text = "SENTINEL AI"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_TEAL
    p0.font.name = "Consolas"

    p1 = tf1.add_paragraph()
    p1.text = "Automated Public Safety & Threat Detection System"
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_HEADER
    p1.font.name = "Arial"
    p1.space_before = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Real-Time Unattended Luggage Detection & Loitering Analytics Engine"
    p2.font.size = Pt(15)
    p2.font.color.rgb = COLOR_MUTED
    p2.font.name = "Arial"
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Final Year Academic Project Evaluation"
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = COLOR_BLUE
    p3.font.name = "Consolas"
    p3.space_before = Pt(25)

    # =========================================================================
    # SLIDE 2: Problem Statement & Objectives
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Problem Statement & Project Objectives")

    # Left Card: Problem Statement
    card1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3))
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_CARD
    card1.line.color.rgb = COLOR_BORDER

    tf1 = card1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "CHALLENGES IN MANUAL SURVEILLANCE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    bullet1 = [
        ("Operator Attention Decay: ", "CCTV operators lose up to 95% of visual attention after 20 minutes of continuous screen monitoring."),
        ("High-Threat Blindspots: ", "Unattended baggage in high-footfall areas (airports, stations) goes unnoticed until critical delay."),
        ("False Alarm Overload: ", "Traditional motion detectors produce constant false alarms without owner-presence context.")
    ]
    for title, desc in bullet1:
        p = tf1.add_paragraph()
        p.space_before = Pt(12)
        run_t = p.add_run()
        run_t.text = title
        run_t.font.bold = True
        run_t.font.size = Pt(11)
        run_t.font.color.rgb = COLOR_HEADER
        run_d = p.add_run()
        run_d.text = desc
        run_d.font.size = Pt(11)
        run_d.font.color.rgb = COLOR_BODY

    # Right Card: Project Objectives
    card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3))
    card2.fill.solid()
    card2.fill.fore_color.rgb = COLOR_CARD
    card2.line.color.rgb = COLOR_BORDER

    tf2 = card2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "PROJECT OBJECTIVES"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    bullet2 = [
        ("Automated Object Detection: ", "Deploy deep learning models to automatically detect luggage and persons simultaneously."),
        ("Temporal Dwell-Time Tracking: ", "Track dwell times and calculate real-time unattended risk scores (0–100%)."),
        ("Smart Owner-Return Cancellation: ", "Disarm alerts automatically when owners return within proximity, reducing false alarms."),
        ("Auditable Reporting Engine: ", "Generate instant PDF incident reports and structured Excel logs for security personnel.")
    ]
    for title, desc in bullet2:
        p = tf2.add_paragraph()
        p.space_before = Pt(10)
        run_t = p.add_run()
        run_t.text = title
        run_t.font.bold = True
        run_t.font.size = Pt(11)
        run_t.font.color.rgb = COLOR_HEADER
        run_d = p.add_run()
        run_d.text = desc
        run_d.font.size = Pt(11)
        run_d.font.color.rgb = COLOR_BODY

    # =========================================================================
    # SLIDE 3: System Architecture Overview
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "System Architecture & Modular Structure")

    modules = [
        ("DUAL INFERENCE ENGINE", "Custom YOLO Luggage Model (best.pt) + Pretrained YOLOv9s Person Model (COCO Class 0)", COLOR_TEAL),
        ("CENTROID TRACKER", "Multi-object ID assignment with Exponential Moving Average (EMA) spatial smoothing", COLOR_BLUE),
        ("RISK ANALYSIS ENGINE", "Dwell time counter, 4-tier risk severity scoring & Proximity owner-return checker", COLOR_AMBER),
        ("GUI DASHBOARD", "Tkinter desktop HUD, live OpenCV video player, parameter controls & Matplotlib graphs", COLOR_TEAL),
        ("AUDIT & REPORTING", "Automated PDF incident generator (reportlab) & session Excel logger (openpyxl)", COLOR_PURPLE)
    ]

    top_pos = 1.6
    for title, desc, color in modules:
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(11.7), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER

        # Left color bar indicator inside card
        cbar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(top_pos + 0.05), Inches(0.08), Inches(0.85))
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = color
        cbar.line.color.rgb = color

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        p = tf.paragraphs[0]
        run_t = p.add_run()
        run_t.text = title + "  ➔  "
        run_t.font.bold = True
        run_t.font.size = Pt(13)
        run_t.font.color.rgb = color
        run_t.font.name = "Consolas"

        run_d = p.add_run()
        run_d.text = desc
        run_d.font.size = Pt(12)
        run_d.font.color.rgb = COLOR_HEADER
        top_pos += 1.05

    # =========================================================================
    # SLIDE 4: Model Training & Evaluation Metrics (NEW - EMPIRICAL DATA)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Model Training & Evaluation Metrics (Empirical Results)")

    # 4 Key Stat Metric Cards at the top
    stats = [
        ("93.07%", "mAP@50", "Detection Accuracy", COLOR_TEAL),
        ("94.96%", "PRECISION", "Low False Positives", COLOR_BLUE),
        ("84.58%", "RECALL", "High Threat Capture", COLOR_PURPLE),
        ("72.86%", "mAP@50-95", "Strict IOU Metric", COLOR_AMBER)
    ]
    s_left = 0.8
    for val, label, subtext, color in stats:
        scard = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(s_left), Inches(1.6), Inches(2.75), Inches(1.7))
        scard.fill.solid()
        scard.fill.fore_color.rgb = COLOR_CARD
        scard.line.color.rgb = COLOR_BORDER

        # Top line accent
        bar = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(s_left + 0.1), Inches(1.65), Inches(2.55), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.color.rgb = color

        tf = scard.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.15)
        
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = val
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Consolas"

        p1 = tf.add_paragraph()
        p1.alignment = PP_ALIGN.CENTER
        p1.text = label
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_HEADER

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = subtext
        p2.font.size = Pt(9)
        p2.font.color.rgb = COLOR_MUTED

        s_left += 2.98

    # Bottom Left Card: Training Config & Loss Progression
    b_card1 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.5), Inches(5.6), Inches(3.4))
    b_card1.fill.solid()
    b_card1.fill.fore_color.rgb = COLOR_CARD
    b_card1.line.color.rgb = COLOR_BORDER

    tf = b_card1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "TRAINING HYPERPARAMETERS & LOSS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    train_info = [
        ("Base Architecture: ", "YOLOv9 Custom Luggage Model"),
        ("Training Duration: ", "30 Epochs (Image size: 640x640, Batch: 4)"),
        ("Final Training Box Loss: ", "0.5878 (Converged cleanly)"),
        ("Final Validation Box Loss: ", "0.7708"),
        ("Final Class Loss: ", "0.3621 (train) / 0.5276 (val)")
    ]
    for t, d in train_info:
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        r1.text = t
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_HEADER
        r2 = p.add_run()
        r2.text = d
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_BODY

    # Bottom Right Card: Automated Excel Training Logger
    b_card2 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.5), Inches(5.7), Inches(3.4))
    b_card2.fill.solid()
    b_card2.fill.fore_color.rgb = COLOR_CARD
    b_card2.line.color.rgb = COLOR_BORDER

    tf = b_card2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "AUTOMATED EXCEL METRIC LOGGING (start_training.py)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    log_info = [
        ("Epoch-by-Epoch Excel Tracking: ", "Custom logging hook automatically writes Box Loss, mAP@50, and timestamps to XLSX per epoch."),
        ("Training Log File: ", "train_training_log.xlsx generated in models/ directory."),
        ("Automated Visualization: ", "Generates BoxPR_curve.png, confusion_matrix.png, and results.png automatically upon completion.")
    ]
    for t, d in log_info:
        p = tf.add_paragraph()
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = t
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_HEADER
        r2 = p.add_run()
        r2.text = d
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_BODY

    # =========================================================================
    # SLIDE 5: End-to-End Ingestion & Processing Pipeline
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Ingestion & Processing Pipeline (5 Phases)")

    phases = [
        ("Phase 1", "Ingestion", "OpenCV reads webcam, RTSP stream, or video file"),
        ("Phase 2", "Dual Inference", "Passes frame to Bag & Person YOLO detectors"),
        ("Phase 3", "Centroid Tracking", "Assigns persistent IDs & applies EMA smoothing"),
        ("Phase 4", "Risk Analysis", "Calculates dwell time & checks owner proximity"),
        ("Phase 5", "HUD & Reporting", "Draws dynamic overlay, triggers alarms & PDF log")
    ]

    left_pos = 0.8
    for p_num, p_title, p_desc in phases:
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(1.8), Inches(2.15), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER

        top_line = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left_pos + 0.1), Inches(1.85), Inches(1.95), Inches(0.06))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = COLOR_TEAL
        top_line.line.color.rgb = COLOR_TEAL

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = p_num.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEAL
        p.font.name = "Consolas"

        p_t = tf.add_paragraph()
        p_t.text = p_title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_HEADER
        p_t.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = p_desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_BODY
        p_d.space_before = Pt(12)

        left_pos += 2.4

    # =========================================================================
    # SLIDE 6: Core AI Algorithms & Innovation
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Core Algorithms & Key Technical Innovations")

    # Card 1: Owner Return Verification
    c1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_CARD
    c1.line.color.rgb = COLOR_BORDER

    tf = c1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. PROXIMITY OWNER-RETURN ALGORITHM"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    items1 = [
        ("Spatial Distance Check: ", "Calculates Euclidean distance between bag centroid (bcx, bcy) and all person centroids (pcx, pcy)."),
        ("Proximity Threshold: ", "If dist ≤ 150 pixels, the system identifies owner presence."),
        ("State Reset: ", "Disarms alert instantly, resets first_seen timer, and displays 'OWNER RETURNED' for 90 frames."),
        ("Impact: ", "Drastically reduces false alarms without operator manual intervention.")
    ]
    for t, d in items1:
        p = tf.add_paragraph()
        p.space_before = Pt(10)
        r1 = p.add_run()
        r1.text = t
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_HEADER
        r2 = p.add_run()
        r2.text = d
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_BODY

    # Card 2: Risk Scoring & Smooth Tracking
    c2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_CARD
    c2.line.color.rgb = COLOR_BORDER

    tf = c2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. RISK SCORING & SMOOTH TRACKING"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    items2 = [
        ("Continuous Risk Equation: ", "Risk Score = min((Elapsed Time / Alert Threshold) × 100, 100)%"),
        ("4-Tier Severity Scale: ", "LOW (0-33%), MEDIUM (34-66%), HIGH (67-99%), CRITICAL (100%)."),
        ("EMA Bounding Box Smoothing: ", "smooth_x = 0.75 * new_x + 0.25 * old_x (Eliminates frame jitter)."),
        ("Occlusion Buffer: ", "Retains tracks for up to 3.0s if object is momentarily blocked.")
    ]
    for t, d in items2:
        p = tf.add_paragraph()
        p.space_before = Pt(10)
        r1 = p.add_run()
        r1.text = t
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_HEADER
        r2 = p.add_run()
        r2.text = d
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_BODY

    # =========================================================================
    # SLIDE 7: Dynamic Dashboard & User Experience
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Desktop Application & User Dashboard")

    features = [
        ("Live OpenCV Video HUD", "Overlays color-coded bounding boxes, risk badges, dwell timers, and alert banners in real-time.", COLOR_TEAL),
        ("Real-Time Parameter Controls", "Interactive sliders allow operators to adjust Alert Thresholds (sec), Loitering Limits, and Confidence Thresholds dynamically.", COLOR_BLUE),
        ("Embedded Matplotlib Analytics", "Displays real-time threat trend graphs (Detections over Time) embedded directly into the Tkinter GUI context.", COLOR_AMBER),
        ("Audio-Visual Alert System", "Combines visual HUD flashing with automatic Windows audio siren alerts (winsound integration).", COLOR_PURPLE)
    ]

    top_pos = 1.6
    for title, desc, color in features:
        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(11.7), Inches(1.1))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER

        cbar = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(top_pos + 0.05), Inches(0.08), Inches(1.0))
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = color
        cbar.line.color.rgb = color

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        p = tf.paragraphs[0]
        r_t = p.add_run()
        r_t.text = title + "\n"
        r_t.font.bold = True
        r_t.font.size = Pt(14)
        r_t.font.color.rgb = color

        r_d = p.add_run()
        r_d.text = desc
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = COLOR_BODY
        top_pos += 1.25

    # =========================================================================
    # SLIDE 8: Audit Logging & Reporting Subsystem
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Reporting, Audit Logging & Headless Testing")

    # Card 1: PDF Incident Reports
    c1 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_CARD
    c1.line.color.rgb = COLOR_BORDER
    tf = c1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PDF INCIDENT REPORTS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL
    p1 = tf.add_paragraph()
    p1.text = "• Generated automatically via reportlab\n• Includes incident timestamp, risk severity, duration, and frame snapshots\n• Official documentation for law enforcement and station security."
    p1.font.size = Pt(11)
    p1.font.color.rgb = COLOR_BODY
    p1.space_before = Pt(12)

    # Card 2: Excel Session Audit Logs
    c2 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_CARD
    c2.line.color.rgb = COLOR_BORDER
    tf = c2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "EXCEL SESSION LOGS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    p2 = tf.add_paragraph()
    p2.text = "• Built with openpyxl\n• Tracks total bag detections, loitering counts, confidence values, and timestamped events\n• Enables post-session statistical analysis."
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_BODY
    p2.space_before = Pt(12)

    # Card 3: Headless Engine
    c3 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2))
    c3.fill.solid()
    c3.fill.fore_color.rgb = COLOR_CARD
    c3.line.color.rgb = COLOR_BORDER
    tf = c3.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "HEADLESS ENGINE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER
    p3 = tf.add_paragraph()
    p3.text = "• test_detection.py script\n• Simulates processing pipeline without GUI rendering\n• Allows benchmark testing of detection accuracy & frame rates on test videos."
    p3.font.size = Pt(11)
    p3.font.color.rgb = COLOR_BODY
    p3.space_before = Pt(12)

    # =========================================================================
    # SLIDE 9: Comparative Advantage Table
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Comparative Advantage & Performance Matrix")

    table_shape = slide9.shapes.add_table(5, 3, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    table = table_shape.table

    headers = ["FEATURE / METRIC", "CONVENTIONAL CCTV / MOTION", "SENTINEL AI SYSTEM"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(241, 245, 249) # Light slate header fill
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEAL if i == 2 else COLOR_HEADER

    rows_data = [
        ("Surveillance Model", "Manual Human Operator (Fatigue-prone)", "100% Automated 24/7 Deep Learning Pipeline"),
        ("Contextual Awareness", "Simple Pixel Motion (High False Alarms)", "Proximity Owner-Return Check (Smart Disarm)"),
        ("Risk Prioritization", "Binary Alert (Motion / No Motion)", "Dynamic 4-Tier Risk Score (0-100% Severity)"),
        ("Deployment & Cost", "High Cloud Streaming Costs", "Edge-Compatible Zero-Cloud Local Execution")
    ]

    for row_idx, row in enumerate(rows_data, 1):
        for col_idx, val in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.bold = (col_idx == 2)
            p.font.color.rgb = COLOR_TEAL if col_idx == 2 else COLOR_BODY

    # =========================================================================
    # SLIDE 10: Future Enhancements
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "Future Enhancements & Scalability Roadmap")

    enhancements = [
        ("Multi-Camera Person Re-Identification (Re-ID)", "Deploy deep feature embedding networks to track owners across non-overlapping camera feeds in large airports or station terminals.", COLOR_TEAL),
        ("TensorRT & GPU Multi-Stream Scaling", "Export PyTorch models to TensorRT ONNX engines to achieve 60+ FPS processing across 16+ simultaneous 4K camera streams.", COLOR_BLUE),
        ("Cloud & Mobile Security Push Notifications", "Integrate MQTT / WebSockets gateway to push instant critical risk alerts and video snapshots directly to security officers' mobile apps.", COLOR_AMBER),
        ("Automated PTZ Camera Tracking", "Send pan-tilt-zoom motor commands to cameras to automatically zoom in and isolate identified high-risk unattended bags.", COLOR_PURPLE)
    ]

    top_pos = 1.6
    for title, desc, color in enhancements:
        card = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(11.7), Inches(1.1))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER

        cbar = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(top_pos + 0.05), Inches(0.08), Inches(1.0))
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = color
        cbar.line.color.rgb = color

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        p = tf.paragraphs[0]
        r_t = p.add_run()
        r_t.text = title + "\n"
        r_t.font.bold = True
        r_t.font.size = Pt(13)
        r_t.font.color.rgb = color

        r_d = p.add_run()
        r_d.text = desc
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = COLOR_BODY
        top_pos += 1.25

    # =========================================================================
    # SLIDE 11: Conclusion & Q&A
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11)

    c_box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = COLOR_CARD
    c_box.line.color.rgb = COLOR_BORDER

    # Top line accent
    top_line = slide11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.6), Inches(1.55), Inches(10.133), Inches(0.08))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = COLOR_TEAL
    top_line.line.color.rgb = COLOR_TEAL

    tf = c_box.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    p0.text = "THANK YOU!"
    p0.font.size = Pt(40)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_TEAL
    p0.font.name = "Consolas"

    p1 = tf.add_paragraph()
    p1.alignment = PP_ALIGN.CENTER
    p1.text = "SENTINEL AI — Public Safety Surveillance & Automated Threat Detection System"
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_HEADER
    p1.space_before = Pt(15)

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Open for Questions & Academic Evaluation Feedback"
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLOR_MUTED
    p2.space_before = Pt(25)

    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SENTINEL_AI_Presentation_Light.pptx")
    prs.save(out_path)
    print(f"Presentation saved successfully to {out_path}")
    
    # Try overwriting original if not locked
    try:
        orig_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SENTINEL_AI_Presentation.pptx")
        prs.save(orig_path)
    except Exception:
        pass


if __name__ == "__main__":
    create_deck()

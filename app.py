"""
SENTINEL AI — Public Safety Surveillance System
Full Tkinter desktop application for AI-based person detection, tracking, behavior, and posture analysis.
"""

# ═══════════════════════════════════════════════════════════════════════════════
# IMPORTS
# ═══════════════════════════════════════════════════════════════════════════════

import tkinter as tk
from tkinter import ttk, filedialog
import threading
import os
import sys
import time
import re
import subprocess
import math

import yaml
import cv2
import numpy as np
from PIL import Image, ImageTk

import matplotlib
matplotlib.use("TkAgg")
from matplotlib.figure import Figure
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

try:
    import winsound
    HAS_WINSOUND = True
except ImportError:
    HAS_WINSOUND = False

# ═══════════════════════════════════════════════════════════════════════════════
# GLOBAL STYLE & CONFIGURATION CONSTANTS
# ═══════════════════════════════════════════════════════════════════════════════

BG       = "#0f141b"
PANEL    = "#151b23"
CARD     = "#1b2430"
CARD_ALT = "#202a36"
ACCENT   = "#5fd38d"
DANGER   = "#ef6b73"
WARNING  = "#d9a441"
INFO     = "#7db7e8"
TEXT     = "#e6edf3"
MUTED    = "#8b98a8"
BORDER   = "#2c3745"

FONT_HEADING = ("Segoe UI", 16, "bold")
FONT_BODY    = ("Segoe UI", 10)
FONT_SMALL   = ("Segoe UI", 9)
FONT_MONO    = ("Consolas", 9)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))


# ═══════════════════════════════════════════════════════════════════════════════
# MODULAR SENTINEL ENGINE INTEGRATION
# ═══════════════════════════════════════════════════════════════════════════════

from sentinel.config import SentinelConfig
from sentinel.detection import (
    OBJECT_GROUP_MAP, get_object_group, is_person,
    load_person_model, detect_persons, calculate_iou, iou_deduplicate
)
from sentinel.tracking import (
    TrackStatus, TrackState, match_entities, match_person_centroids
)
from sentinel.behavior import (
    compute_dwell_time, compute_trajectory_metrics, update_grid_revisits,
    evaluate_person_behavior
)
from sentinel.risk import calculate_entity_risk, analyze_behaviors
from sentinel.alerting import draw_overlays
from sentinel.posture import load_pose_model, estimate_postures


def frame_to_tk(frame, canvas_w, canvas_h):
    """Convert cv2 BGR frame to ImageTk.PhotoImage sized to fit canvas."""
    if frame is None or not hasattr(frame, "shape") or frame.size == 0:
        return None
    h, w = frame.shape[:2]
    if w <= 0 or h <= 0:
        return None
    scale = min(canvas_w / w, canvas_h / h)
    new_w = max(int(w * scale), 1)
    new_h = max(int(h * scale), 1)
    resized = cv2.resize(frame, (new_w, new_h), interpolation=cv2.INTER_AREA)
    rgb = cv2.cvtColor(resized, cv2.COLOR_BGR2RGB)
    return ImageTk.PhotoImage(Image.fromarray(rgb))


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN APPLICATION CLASS
# ═══════════════════════════════════════════════════════════════════════════════


class SentinelApp:
    def __init__(self, root):
        self.root = root
        self.root.title("SENTINEL AI \u2014 Public Safety Surveillance")
        self.root.geometry("1400x900")
        self.root.resizable(True, True)
        self.root.configure(bg=BG)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

        # State
        self.status_var = tk.StringVar(value="IDLE")
        self.training_process = None
        self.video_running = False
        self.camera_running = False
        self.current_training_epoch = 0
        self.train_losses = []
        self.map_values = []
        self.loss_epochs = []
        self.map_epochs = []
        self._last_frame = None
        self._alert_sound_active = False
        self._alert_sound_thread = None
        self.app_pages = {}
        self.nav_buttons = {}
        self.current_page = None
        self.incidents = []
        self.unread_incidents = 0
        self.vid_processing_fps = 0.0
        self.cam_processing_fps = 0.0
        self.video_duration = 0.0
        self.video_position = 0.0
        self.video_source_status = "No file selected"
        self.camera_source_status = "Disconnected"
        self.model_status_text = "Models not loaded"
        self.pose_status_text = "Pose model not loaded"

        self._setup_styles()
        self._build_header()
        self._build_notebook()
        self._build_status_bar()
        self._build_dataset_tab()
        self._build_train_tab()
        self._build_video_tab()
        self._build_camera_tab()
        self._build_incidents_page()
        self._build_overview_page()
        self._build_settings_page()
        self._show_page("overview")
        self._refresh_overview()
        self._update_clock()

    # ── Cleanup ──────────────────────────────────────────────────────────────

    def _on_close(self):
        self.video_running = False
        self.camera_running = False
        if self.training_process:
            try:
                self.training_process.terminate()
            except Exception:
                pass
        self.root.destroy()

    # ── Styles ───────────────────────────────────────────────────────────────

    def _setup_styles(self):
        style = ttk.Style()
        style.theme_use("default")

        style.configure("TFrame", background=BG)
        style.configure("TLabel", background=BG, foreground=TEXT, font=FONT_BODY)
        style.configure("TCombobox", fieldbackground=CARD, background=CARD,
                        foreground=TEXT, arrowcolor=TEXT, bordercolor=BORDER)
        style.map("TCombobox", fieldbackground=[("readonly", CARD)],
                  foreground=[("readonly", TEXT)])
        style.configure("Treeview", background=CARD, fieldbackground=CARD,
                        foreground=TEXT, rowheight=26, borderwidth=0,
                        font=FONT_SMALL)
        style.configure("Treeview.Heading", background=PANEL, foreground=TEXT,
                        font=("Segoe UI", 9, "bold"), borderwidth=0)
        style.map("Treeview", background=[("selected", CARD_ALT)],
                  foreground=[("selected", TEXT)])

        style.configure("Horizontal.TProgressbar",
                         background=ACCENT, troughcolor=CARD, borderwidth=0, thickness=6)

    # ── Header ───────────────────────────────────────────────────────────────

    def _build_header(self):
        header = tk.Frame(self.root, bg=PANEL, height=60)
        header.pack(fill="x", side="top")
        header.pack_propagate(False)

        tk.Label(header, text="\u2B21 SENTINEL AI", font=FONT_HEADING,
                 fg=ACCENT, bg=PANEL).pack(side="left", padx=20)

        tk.Label(header, text="PUBLIC SAFETY SURVEILLANCE", font=FONT_SMALL,
                 fg=MUTED, bg=PANEL).pack(side="left", padx=(0, 20))

        self.clock_label = tk.Label(header, text="", font=FONT_BODY, fg=MUTED, bg=PANEL)
        self.clock_label.pack(side="right", padx=20)

    def _update_clock(self):
        self.clock_label.config(text=time.strftime("%Y-%m-%d  %H:%M:%S"))
        self.root.after(1000, self._update_clock)

    # ── Notebook ─────────────────────────────────────────────────────────────

    def _build_notebook(self):
        shell = tk.Frame(self.root, bg=BG)
        shell.pack(fill="both", expand=True)

        self.sidebar = tk.Frame(shell, bg=PANEL, width=210)
        self.sidebar.pack(side="left", fill="y")
        self.sidebar.pack_propagate(False)

        brand = tk.Frame(self.sidebar, bg=PANEL)
        brand.pack(fill="x", padx=14, pady=(18, 14))
        tk.Label(brand, text="OPERATIONS", font=("Segoe UI", 9, "bold"),
                 fg=MUTED, bg=PANEL).pack(anchor="w")
        self.nav_incident_badge = tk.Label(brand, text="0 new incidents",
                                           font=FONT_SMALL, fg=MUTED, bg=PANEL)
        self.nav_incident_badge.pack(anchor="w", pady=(3, 0))

        self.content = tk.Frame(shell, bg=BG)
        self.content.pack(side="left", fill="both", expand=True)
        self.content.grid_rowconfigure(0, weight=1)
        self.content.grid_columnconfigure(0, weight=1)

        nav_items = [
            ("overview", "Overview"),
            ("camera", "Live Monitoring"),
            ("video", "Video Analysis"),
            ("incidents", "Incidents"),
            ("dataset", "Dataset"),
            ("train", "Training"),
            ("stats", "Statistics"),
            ("settings", "Settings"),
        ]
        for key, label in nav_items:
            btn = tk.Button(self.sidebar, text=label, anchor="w", font=FONT_BODY,
                            bg=PANEL, fg=TEXT, activebackground=CARD_ALT,
                            activeforeground=TEXT, relief="flat", bd=0,
                            cursor="hand2", padx=18, pady=11,
                            command=lambda k=key: self._show_page(k))
            btn.pack(fill="x", padx=10, pady=2)
            self.nav_buttons[key] = btn

        self.tab_overview = self._new_page("overview")
        self.tab_camera = self._new_page("camera")
        self.tab_video = self._new_page("video")
        self.tab_incidents = self._new_page("incidents")
        self.tab_dataset = self._new_page("dataset")
        self.tab_train = self._new_page("train")
        self.tab_stats = self._new_page("stats")
        self.tab_settings = self._new_page("settings")
        self._build_stats_tab()

    def _new_page(self, name):
        page = tk.Frame(self.content, bg=BG)
        page.grid(row=0, column=0, sticky="nsew")
        self.app_pages[name] = page
        return page

    def _show_page(self, name):
        page = self.app_pages.get(name)
        if not page:
            return
        page.tkraise()
        self.current_page = name
        for key, btn in self.nav_buttons.items():
            selected = key == name
            btn.config(bg=CARD_ALT if selected else PANEL,
                       fg=ACCENT if selected else TEXT)
        if name == "incidents":
            self.unread_incidents = 0
            self._update_incident_badges()

    # ── Status Bar ───────────────────────────────────────────────────────────

    def _build_status_bar(self):
        bar = tk.Frame(self.root, bg=PANEL, height=28)
        bar.pack(fill="x", side="bottom")
        bar.pack_propagate(False)

        tk.Label(bar, text="\u25CF", font=("Consolas", 8), fg=ACCENT,
                 bg=PANEL).pack(side="left", padx=(15, 4))

        self.status_label = tk.Label(bar, textvariable=self.status_var,
                                      font=FONT_SMALL, fg=ACCENT, bg=PANEL)
        self.status_label.pack(side="left")

    def _set_status(self, text, color=ACCENT):
        self.status_var.set(text)
        self.status_label.config(fg=color)

    # ── Shared UI builders ─────────────────────────────────────────────────

    def _section_header(self, parent, title, subtitle=None):
        box = tk.Frame(parent, bg=BG)
        box.pack(fill="x", padx=18, pady=(16, 8))
        tk.Label(box, text=title, font=("Segoe UI", 18, "bold"),
                 fg=TEXT, bg=BG).pack(anchor="w")
        if subtitle:
            tk.Label(box, text=subtitle, font=FONT_SMALL,
                     fg=MUTED, bg=BG).pack(anchor="w", pady=(2, 0))
        return box

    def _metric_card(self, parent, title, value_var, accent=ACCENT):
        card = tk.Frame(parent, bg=CARD, highlightthickness=1, highlightbackground=BORDER)
        tk.Label(card, text=title, font=("Segoe UI", 9, "bold"),
                 fg=MUTED, bg=CARD).pack(anchor="w", padx=14, pady=(12, 0))
        tk.Label(card, textvariable=value_var, font=("Segoe UI", 19, "bold"),
                 fg=accent, bg=CARD).pack(anchor="w", padx=14, pady=(4, 12))
        return card

    def _set_text(self, widget, text):
        widget.config(state="normal")
        widget.delete("1.0", "end")
        widget.insert("1.0", text)
        widget.config(state="disabled")

    # ═════════════════════════════════════════════════════════════════════════
    # PAGE — OVERVIEW
    # ═════════════════════════════════════════════════════════════════════════

    def _build_overview_page(self):
        self._section_header(
            self.tab_overview,
            "Overview",
            "Live operational state from the current camera, video analysis, models, and incident stream."
        )

        self.overview_vars = {
            "camera": tk.StringVar(value="Disconnected"),
            "fps": tk.StringVar(value="0.0 FPS"),
            "people": tk.StringVar(value="0"),
            "incidents": tk.StringVar(value="0 active"),
            "models": tk.StringVar(value="Not loaded"),
        }

        cards = tk.Frame(self.tab_overview, bg=BG)
        cards.pack(fill="x", padx=18, pady=(4, 12))
        items = [
            ("Camera / Stream", "camera", INFO),
            ("Processing FPS", "fps", ACCENT),
            ("Tracked People", "people", INFO),
            ("Active Incidents", "incidents", DANGER),
            ("Model Status", "models", ACCENT),
        ]
        for idx, (title, key, color) in enumerate(items):
            card = self._metric_card(cards, title, self.overview_vars[key], color)
            card.grid(row=idx // 3, column=idx % 3, sticky="nsew", padx=6, pady=6)
        for col in range(3):
            cards.grid_columnconfigure(col, weight=1)

        body = tk.Frame(self.tab_overview, bg=BG)
        body.pack(fill="both", expand=True, padx=18, pady=(0, 18))
        body.grid_columnconfigure(0, weight=3)
        body.grid_columnconfigure(1, weight=2)
        body.grid_rowconfigure(0, weight=1)

        events_panel = tk.Frame(body, bg=CARD, highlightthickness=1, highlightbackground=BORDER)
        events_panel.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        tk.Label(events_panel, text="Recent Events", font=("Segoe UI", 12, "bold"),
                 fg=TEXT, bg=CARD).pack(anchor="w", padx=14, pady=(12, 4))
        self.overview_events = tk.Text(events_panel, bg=CARD, fg=TEXT, font=FONT_MONO,
                                       relief="flat", padx=12, pady=8, height=12,
                                       highlightthickness=0)
        self.overview_events.pack(fill="both", expand=True, padx=8, pady=(0, 10))
        self.overview_events.config(state="disabled")

        model_panel = tk.Frame(body, bg=CARD, highlightthickness=1, highlightbackground=BORDER)
        model_panel.grid(row=0, column=1, sticky="nsew", padx=(8, 0))
        tk.Label(model_panel, text="Runtime", font=("Segoe UI", 12, "bold"),
                 fg=TEXT, bg=CARD).pack(anchor="w", padx=14, pady=(12, 4))
        self.overview_runtime = tk.Text(model_panel, bg=CARD, fg=TEXT, font=FONT_MONO,
                                        relief="flat", padx=12, pady=8,
                                        highlightthickness=0)
        self.overview_runtime.pack(fill="both", expand=True, padx=8, pady=(0, 10))
        self.overview_runtime.config(state="disabled")

    def _refresh_overview(self):
        cam_entities = getattr(self, "cam_tracked_entities", {})
        vid_entities = getattr(self, "vid_tracked_entities", {})
        active_entities = cam_entities if self.camera_running else vid_entities
        people = sum(1 for e in active_entities.values() if is_person(e.get("class_name", "")))
        active_incidents = sum(1 for i in self.incidents if not i.get("acknowledged"))
        fps = self.cam_processing_fps if self.camera_running else self.vid_processing_fps
        stream = self.camera_source_status if self.camera_running else self.video_source_status
        if not self.camera_running and not self.video_running and self.video_path is None:
            stream = self.camera_source_status

        self.overview_vars["camera"].set(stream)
        self.overview_vars["fps"].set(f"{fps:.1f} FPS")
        self.overview_vars["people"].set(str(people))
        self.overview_vars["incidents"].set(f"{active_incidents} active")
        self.overview_vars["models"].set(self.model_status_text)

        recent = self.incidents[-10:]
        if recent:
            lines = [
                f"[{i['time']}] {i['severity']:<8} {i['category']:<22} "
                f"{i['object_label']}  {i['summary']}"
                for i in reversed(recent)
            ]
            self._set_text(self.overview_events, "\n".join(lines))
        else:
            self._set_text(self.overview_events, "No incidents recorded this session.")

        runtime = (
            f"Application status: {self.status_var.get()}\n"
            f"Camera: {self.camera_source_status}\n"
            f"Video: {self.video_source_status}\n"
            f"Video time: {self._format_seconds(self.video_position)} / {self._format_seconds(self.video_duration)}\n"
            f"Detection models: {self.model_status_text}\n"
            f"Pose model: {self.pose_status_text}\n"
            f"Confidence threshold: camera {getattr(self, 'cam_conf', tk.DoubleVar(value=0.25)).get():.2f}, "
            f"video {getattr(self, 'vid_conf', tk.DoubleVar(value=0.25)).get():.2f}"
        )
        self._set_text(self.overview_runtime, runtime)
        self.root.after(1000, self._refresh_overview)

    # ═════════════════════════════════════════════════════════════════════════
    # PAGE — INCIDENTS
    # ═════════════════════════════════════════════════════════════════════════

    def _build_incidents_page(self):
        self._section_header(
            self.tab_incidents,
            "Incidents",
            "Non-modal incident history with severity, evidence, acknowledgement and snapshot paths."
        )
        filters = tk.Frame(self.tab_incidents, bg=BG)
        filters.pack(fill="x", padx=18, pady=(0, 8))
        tk.Label(filters, text="Category", font=FONT_SMALL, fg=MUTED, bg=BG).pack(side="left")
        self.incident_category_filter = tk.StringVar(value="All")
        self.incident_category_combo = ttk.Combobox(filters, textvariable=self.incident_category_filter,
                                                    values=["All"], state="readonly", width=22)
        self.incident_category_combo.pack(side="left", padx=(6, 16))
        tk.Label(filters, text="Severity", font=FONT_SMALL, fg=MUTED, bg=BG).pack(side="left")
        self.incident_severity_filter = tk.StringVar(value="All")
        ttk.Combobox(filters, textvariable=self.incident_severity_filter,
                     values=["All", "LOW", "MEDIUM", "HIGH", "CRITICAL"], state="readonly",
                     width=12).pack(side="left", padx=(6, 16))
        tk.Button(filters, text="Apply Filter", font=FONT_SMALL, bg=CARD, fg=TEXT,
                  relief="flat", command=self._refresh_incidents_view).pack(side="left")
        tk.Button(filters, text="Acknowledge Selected", font=FONT_SMALL, bg=CARD, fg=ACCENT,
                  relief="flat", command=self._ack_selected_incident).pack(side="left", padx=8)
        tk.Button(filters, text="Dismiss Selected", font=FONT_SMALL, bg=CARD, fg=WARNING,
                  relief="flat", command=self._dismiss_selected_incident).pack(side="left")

        columns = ("time", "source", "category", "severity", "target", "status", "evidence")
        self.incident_tree = ttk.Treeview(self.tab_incidents, columns=columns, show="headings", height=14)
        widths = [90, 90, 170, 90, 90, 120, 440]
        for col, width in zip(columns, widths):
            self.incident_tree.heading(col, text=col.title())
            self.incident_tree.column(col, width=width, anchor="w")
        self.incident_tree.pack(fill="both", expand=True, padx=18, pady=(0, 8))
        self.incident_tree.bind("<<TreeviewSelect>>", lambda _e: self._show_selected_incident_details())

        details = tk.Frame(self.tab_incidents, bg=CARD, highlightthickness=1, highlightbackground=BORDER, height=130)
        details.pack(fill="x", padx=18, pady=(0, 18))
        details.pack_propagate(False)
        tk.Label(details, text="Incident Details", font=("Segoe UI", 11, "bold"),
                 fg=TEXT, bg=CARD).pack(anchor="w", padx=12, pady=(8, 0))
        self.incident_details = tk.Text(details, bg=CARD, fg=TEXT, font=FONT_MONO,
                                        relief="flat", height=5, padx=10, pady=6,
                                        highlightthickness=0)
        self.incident_details.pack(fill="both", expand=True, padx=8, pady=(0, 8))
        self.incident_details.config(state="disabled")

    def _build_settings_page(self):
        self._section_header(
            self.tab_settings,
            "Settings",
            "Runtime configuration that is already supported by the current application."
        )
        panel = tk.Frame(self.tab_settings, bg=CARD, highlightthickness=1, highlightbackground=BORDER)
        panel.pack(fill="x", padx=18, pady=8)
        self.settings_text = tk.Text(panel, bg=CARD, fg=TEXT, font=FONT_MONO,
                                     relief="flat", padx=12, pady=12,
                                     height=14, highlightthickness=0)
        self.settings_text.pack(fill="x")
        self.settings_text.config(state="disabled")
        tk.Button(self.tab_settings, text="Refresh Runtime Settings", font=FONT_BODY,
                  bg=CARD, fg=ACCENT, relief="flat",
                  command=self._refresh_settings_page).pack(anchor="w", padx=18, pady=8)
        self._refresh_settings_page()

    def _refresh_settings_page(self):
        text = (
            f"Person confidence default: {SentinelConfig.PERSON_CONF_DEFAULT:.2f}\n"
            f"Camera confidence slider: {getattr(self, 'cam_conf', tk.DoubleVar(value=0.25)).get():.2f}\n"
            f"Video confidence slider: {getattr(self, 'vid_conf', tk.DoubleVar(value=0.25)).get():.2f}\n"
            f"Camera index: {getattr(self, 'cam_idx_var', tk.StringVar(value='0')).get()}\n"
            f"Dwell threshold: {SentinelConfig.DWELL_TIME_THRESHOLD:.0f}s\n"
            f"Fall confirmation frames: {SentinelConfig.POSTURE_FALL_CONFIRMATION_FRAMES}\n"
            f"Alert cooldown: {SentinelConfig.ALERT_COOLDOWN:.0f}s\n"
            f"Detection models: {self.model_status_text}\n"
            f"Pose model: {self.pose_status_text}"
        )
        self._set_text(self.settings_text, text)

    # ═════════════════════════════════════════════════════════════════════════
    # TAB 1 — DATASET
    # ═════════════════════════════════════════════════════════════════════════

    def _build_dataset_tab(self):
        # ── Left panel ──
        left = tk.Frame(self.tab_dataset, bg=PANEL, width=400)
        left.pack(side="left", fill="y", padx=(10, 5), pady=10)
        left.pack_propagate(False)

        tk.Label(left, text="DATASET VERSIONS", font=FONT_HEADING,
                 fg=ACCENT, bg=PANEL).pack(padx=15, pady=(15, 10), anchor="w")

        self.dataset_list = tk.Text(left, bg=CARD, fg=TEXT, font=FONT_BODY,
                                     relief="flat", height=20, width=40,
                                     insertbackground=TEXT, padx=12, pady=10,
                                     highlightthickness=1, highlightbackground=MUTED)
        self.dataset_list.pack(padx=15, pady=5, fill="both", expand=True)
        self._populate_dataset_list()
        self.dataset_list.config(state="disabled")

        self.verify_progress = ttk.Progressbar(left, mode="indeterminate",
                                                 style="Horizontal.TProgressbar")
        self.verify_progress.pack(padx=15, pady=8, fill="x")

        tk.Button(left, text="\u25B6 RUN VERIFICATION", font=FONT_BODY,
                  bg=ACCENT, fg=BG, relief="flat", cursor="hand2",
                  activebackground="#00cc6a", activeforeground=BG,
                  command=self._run_verification).pack(padx=15, pady=(0, 15),
                                                        fill="x", ipady=8)

        # ── Right panel ──
        right = tk.Frame(self.tab_dataset, bg=PANEL)
        right.pack(side="right", fill="both", expand=True, padx=(5, 10), pady=10)

        tk.Label(right, text="VERIFICATION RESULTS", font=FONT_HEADING,
                 fg=ACCENT, bg=PANEL).pack(padx=15, pady=(15, 10), anchor="w")

        self.verify_output = tk.Text(right, bg=CARD, fg=TEXT, font=FONT_BODY,
                                      relief="flat", insertbackground=TEXT,
                                      padx=12, pady=10,
                                      highlightthickness=1, highlightbackground=MUTED)
        self.verify_output.pack(padx=15, pady=(5, 15), fill="both", expand=True)
        self.verify_output.tag_config("green",   foreground=ACCENT)
        self.verify_output.tag_config("red",     foreground=DANGER)
        self.verify_output.tag_config("header",  foreground=ACCENT, font=FONT_HEADING)
        self.verify_output.tag_config("warning", foreground=WARNING)
        self.verify_output.config(state="disabled")

    def _populate_dataset_list(self):
        dataset_dir = os.path.join(BASE_DIR, "dataset")
        self.dataset_list.config(state="normal")
        self.dataset_list.delete("1.0", "end")

        if os.path.isdir(dataset_dir):
            folders = sorted(d for d in os.listdir(dataset_dir)
                             if os.path.isdir(os.path.join(dataset_dir, d)))
            if folders:
                for folder in folders:
                    self.dataset_list.insert("end", f"  \U0001F4C1 {folder}\n")
            else:
                self.dataset_list.insert("end", "  No dataset folders found.\n")
        else:
            self.dataset_list.insert("end", "  dataset/ directory not found.\n")

        self.dataset_list.config(state="disabled")

    # ── Verification ─────────────────────────────────────────────────────────

    def _run_verification(self):
        self.verify_output.config(state="normal")
        self.verify_output.delete("1.0", "end")
        self.verify_output.config(state="disabled")
        self.verify_progress.start(12)
        self._set_status("VERIFYING", WARNING)
        threading.Thread(target=self._verification_worker, daemon=True).start()

    def _verification_worker(self):
        try:
            yaml_path = os.path.join(BASE_DIR, "data.yaml")
            if not os.path.isfile(yaml_path):
                self.root.after(0, self._vlog, "\u274C data.yaml not found!\n", "red")
                self.root.after(0, self._verification_done, False)
                return

            with open(yaml_path, "r", encoding="utf-8") as f:
                data = yaml.safe_load(f)

            self.root.after(0, self._vlog,
                            "\u2501\u2501\u2501 DATA.YAML LOADED \u2501\u2501\u2501\n", "header")
            self.root.after(0, self._vlog,
                            f"  Classes: {data.get('names', 'N/A')}\n", "green")
            self.root.after(0, self._vlog,
                            f"  NC:      {data.get('nc', 'N/A')}\n\n", "green")

            issues = False

            for split in ("train", "val", "test"):
                split_path = data.get(split, "")
                if not split_path:
                    self.root.after(0, self._vlog,
                                    f"\u26A0 {split} path not defined in data.yaml\n", "red")
                    issues = True
                    continue

                if not os.path.isabs(split_path):
                    split_path = os.path.normpath(os.path.join(BASE_DIR, split_path))

                images_dir = split_path
                labels_dir = split_path.replace("images", "labels")

                self.root.after(0, self._vlog,
                                f"\u2501\u2501\u2501 {split.upper()} SPLIT \u2501\u2501\u2501\n", "header")

                if not os.path.isdir(images_dir):
                    self.root.after(0, self._vlog,
                                    f"  \u274C Images dir not found:\n  {images_dir}\n", "red")
                    issues = True
                    continue

                imgs = [f for f in os.listdir(images_dir)
                        if f.lower().endswith((".jpg", ".jpeg", ".png"))]
                self.root.after(0, self._vlog,
                                f"  \U0001F4F7 Images: {len(imgs)}\n", "green")

                if os.path.isdir(labels_dir):
                    lbls = [f for f in os.listdir(labels_dir) if f.endswith(".txt")]
                    self.root.after(0, self._vlog,
                                    f"  \U0001F3F7\uFE0F  Labels: {len(lbls)}\n", "green")

                    zero = [f for f in lbls
                            if os.path.getsize(os.path.join(labels_dir, f)) == 0]
                    if zero:
                        issues = True
                        self.root.after(0, self._vlog,
                                        f"  \u26A0 Zero-byte labels: {len(zero)}\n", "red")
                        for zf in zero[:5]:
                            self.root.after(0, self._vlog, f"    \u2192 {zf}\n", "red")
                        if len(zero) > 5:
                            self.root.after(0, self._vlog,
                                            f"    ... and {len(zero)-5} more\n", "red")

                    no_lbl = [img for img in imgs
                              if not os.path.isfile(
                                  os.path.join(labels_dir,
                                               os.path.splitext(img)[0] + ".txt"))]
                    if no_lbl:
                        issues = True
                        self.root.after(0, self._vlog,
                                        f"  \u26A0 Images without labels: {len(no_lbl)}\n", "red")
                        for nl in no_lbl[:5]:
                            self.root.after(0, self._vlog, f"    \u2192 {nl}\n", "red")
                        if len(no_lbl) > 5:
                            self.root.after(0, self._vlog,
                                            f"    ... and {len(no_lbl)-5} more\n", "red")
                else:
                    self.root.after(0, self._vlog,
                                    f"  \u274C Labels dir not found:\n  {labels_dir}\n", "red")
                    issues = True

                self.root.after(0, self._vlog, "\n", None)
                time.sleep(0.3)

            self.root.after(0, self._verification_done, not issues)

        except Exception as e:
            print(f"Verification error: {e}")
            self.root.after(0, self._vlog, f"\u274C Error: {e}\n", "red")
            self.root.after(0, self._verification_done, False)

    def _vlog(self, text, tag):
        self.verify_output.config(state="normal")
        if tag:
            self.verify_output.insert("end", text, tag)
        else:
            self.verify_output.insert("end", text)
        self.verify_output.see("end")
        self.verify_output.config(state="disabled")

    def _verification_done(self, success):
        self.verify_progress.stop()
        if success:
            self._vlog("\n\u2705 DATASET READY \u2014 proceed to training\n", "green")
        else:
            self._vlog("\n\u26A0 ISSUES FOUND \u2014 review flagged files\n", "red")
        self._set_status("IDLE", ACCENT)

    # ═════════════════════════════════════════════════════════════════════════
    # TAB 2 — TRAIN
    # ═════════════════════════════════════════════════════════════════════════

    def _build_train_tab(self):
        # ── Left controls ──
        left = tk.Frame(self.tab_train, bg=PANEL, width=300)
        left.pack(side="left", fill="y", padx=(10, 5), pady=10)
        left.pack_propagate(False)

        tk.Label(left, text="TRAINING CONFIG", font=FONT_HEADING,
                 fg=ACCENT, bg=PANEL).pack(padx=15, pady=(15, 15), anchor="w")

        def _field(parent, label_text, default):
            tk.Label(parent, text=label_text, font=FONT_SMALL,
                     fg=MUTED, bg=PANEL).pack(padx=15, anchor="w")
            var = tk.StringVar(value=default)
            tk.Entry(parent, textvariable=var, font=FONT_BODY,
                     bg=CARD, fg=TEXT, insertbackground=TEXT,
                     relief="flat", highlightthickness=1,
                     highlightbackground=MUTED).pack(padx=15, pady=(2, 10),
                                                      fill="x", ipady=4)
            return var

        # Model dropdown
        tk.Label(left, text="MODEL", font=FONT_SMALL, fg=MUTED,
                 bg=PANEL).pack(padx=15, anchor="w")
        available_models = sorted(
            os.path.splitext(f)[0]
            for f in os.listdir(BASE_DIR)
            if f.endswith(".pt") and os.path.isfile(os.path.join(BASE_DIR, f))
        )
        default_model = available_models[0] if available_models else "yolov9s"
        self.model_var = tk.StringVar(value=default_model)
        self.model_size_var = self.model_var
        ttk.Combobox(left, textvariable=self.model_var,
                      values=available_models if available_models else ["yolov9s"],
                      state="readonly", font=FONT_BODY
                      ).pack(padx=15, pady=(2, 10), fill="x")

        self.epochs_var  = _field(left, "EPOCHS",         "50")
        self.imgsz_var   = _field(left, "IMAGE SIZE",     "640")
        self.imgsize_var = self.imgsz_var
        self.batch_var   = _field(left, "BATCH SIZE",     "4")
        self.savedir_var = _field(left, "SAVE DIRECTORY", "models/")

        self.btn_train = tk.Button(left, text="\u25B6 START TRAINING", font=FONT_BODY,
                                    bg=ACCENT, fg=BG, relief="flat", cursor="hand2",
                                    activebackground="#00cc6a", activeforeground=BG,
                                    command=self._start_training)
        self.btn_train.pack(padx=15, pady=(8, 4), fill="x", ipady=8)

        self.btn_stop_train = tk.Button(left, text="\u25A0 STOP", font=FONT_BODY,
                                         bg=DANGER, fg=TEXT, relief="flat", cursor="hand2",
                                         activebackground="#cc2929",
                                         command=self._stop_training, state="disabled")
        self.btn_stop_train.pack(padx=15, pady=(0, 8), fill="x", ipady=8)

        self.btn_test = tk.Button(left, text="\u25B6 TEST MODEL", font=FONT_BODY,
                                   bg=CARD, fg=MUTED, relief="flat", cursor="hand2",
                                   command=self._test_model, state="disabled")
        self.btn_test.pack(padx=15, pady=(0, 15), fill="x", ipady=8)

        # ── Right panel ──
        right = tk.Frame(self.tab_train, bg=PANEL)
        right.pack(side="right", fill="both", expand=True, padx=(5, 10), pady=10)

        # Charts
        chart_frame = tk.Frame(right, bg=CARD, height=270)
        chart_frame.pack(padx=10, pady=(10, 5), fill="x")
        chart_frame.pack_propagate(False)

        self.fig = Figure(figsize=(8, 2.8), dpi=100, facecolor=CARD)
        self.fig.subplots_adjust(left=0.07, right=0.96, top=0.87, bottom=0.2, wspace=0.3)
        self.ax_loss = self.fig.add_subplot(121)
        self.ax_map  = self.fig.add_subplot(122)
        self._style_axis(self.ax_loss, "Train Loss", "Epoch", "Loss")
        self._style_axis(self.ax_map,  "mAP@50",     "Epoch", "mAP")

        self.chart_canvas = FigureCanvasTkAgg(self.fig, master=chart_frame)
        self.chart_canvas.get_tk_widget().pack(fill="both", expand=True)
        self.chart_canvas.draw()

        # Training log
        tk.Label(right, text="TRAINING LOG", font=FONT_SMALL, fg=MUTED,
                 bg=PANEL).pack(padx=15, pady=(10, 2), anchor="w")

        log_frame = tk.Frame(right, bg=CARD)
        log_frame.pack(padx=10, pady=(0, 10), fill="both", expand=True)

        scrollbar = tk.Scrollbar(log_frame, bg=CARD, troughcolor=CARD)
        scrollbar.pack(side="right", fill="y")

        self.train_log = tk.Text(log_frame, bg=CARD, fg=TEXT, font=FONT_SMALL,
                                  relief="flat", insertbackground=TEXT,
                                  padx=10, pady=8, yscrollcommand=scrollbar.set,
                                  highlightthickness=0)
        self.train_log.pack(side="left", fill="both", expand=True)
        scrollbar.config(command=self.train_log.yview)
        self.train_log.config(state="disabled")
        self.excel_log_path = None

    def _style_axis(self, ax, title, xlabel, ylabel):
        ax.set_facecolor(CARD)
        ax.set_title(title, color=ACCENT, fontsize=10, fontfamily="Consolas")
        ax.set_xlabel(xlabel, color=MUTED, fontsize=8, fontfamily="Consolas")
        ax.set_ylabel(ylabel, color=MUTED, fontsize=8, fontfamily="Consolas")
        ax.tick_params(colors=MUTED, labelsize=7)
        for spine in ax.spines.values():
            spine.set_color(MUTED)
            spine.set_linewidth(0.5)
        ax.grid(True, alpha=0.15, color=MUTED)

    # ── Training control ─────────────────────────────────────────────────────

    def _start_training(self):
        try:
            epochs_val = int(self.epochs_var.get())
            imgsz_val  = int(self.imgsz_var.get())
            batch_val  = int(self.batch_var.get())
            if epochs_val <= 0 or imgsz_val <= 0 or batch_val <= 0:
                raise ValueError("Values must be positive")
        except ValueError:
            self._set_status("INVALID TRAINING INPUT", DANGER)
            self._tlog("Epochs, Image Size, and Batch Size must be positive integers.")
            return

        self.train_losses.clear()
        self.map_values.clear()
        self.loss_epochs.clear()
        self.map_epochs.clear()
        self.current_training_epoch = 0

        self.btn_train.config(state="disabled")
        self.btn_stop_train.config(state="normal")
        self.btn_test.config(state="disabled", bg=CARD, fg=MUTED)
        self._set_status("TRAINING", WARNING)

        self.train_log.config(state="normal")
        self.train_log.delete("1.0", "end")
        self.train_log.config(state="disabled")

        threading.Thread(target=self._run_training_thread, daemon=True).start()

    def _run_training_thread(self):
        model_size = self.model_size_var.get().strip()
        if not model_size.endswith(".pt"):
            model_size += ".pt"

        epochs = self.epochs_var.get().strip()
        imgsz = self.imgsize_var.get().strip()
        batch = self.batch_var.get().strip()
        project_dir = self.savedir_var.get().strip().replace("\\", "/")
        data_yaml = os.path.abspath(os.path.join(BASE_DIR, "data.yaml")).replace("\\", "/")

        run_name = f"run_{int(time.time())}"
        save_dir = os.path.join(BASE_DIR, self.savedir_var.get().strip())
        self.root.after(0, self._init_excel_log, save_dir, run_name)

        script = (
            "import sys, os; "
            "sys.stdout.reconfigure(encoding='utf-8', line_buffering=True); "
            "from ultralytics import YOLO; "
            f"m = YOLO(r'{model_size}'); "
            f"m.train(data=r'{data_yaml}', epochs={epochs}, imgsz={imgsz}, "
            f"batch={batch}, workers=0, project=r'{project_dir}', name='{run_name}', exist_ok=True, verbose=True)"
        )
        cmd = [sys.executable, "-u", "-c", script]

        env = os.environ.copy()
        env["PYTHONUNBUFFERED"] = "1"

        try:
            proc = subprocess.Popen(
                cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                text=True, encoding='utf-8', errors='replace', bufsize=1,
                cwd=BASE_DIR, env=env
            )
            self.training_process = proc

            epoch_loss_regex = re.compile(r'^\s*(\d+)/(\d+)\s+([\d\.]+G?|0G)\s+([\d\.]+)\s+([\d\.]+)\s+([\d\.]+)')
            map_regex = re.compile(r'^\s*all\s+\d+\s+\d+\s+[\d\.]+\s+[\d\.]+\s+([\d\.]+)')

            last_map_val = 0.0
            if proc.stdout:
                for line in iter(proc.stdout.readline, ''):
                    if not line:
                        break

                    self.root.after(0, self._tlog, line.rstrip('\r\n'))

                    loss_match = epoch_loss_regex.search(line)
                    if loss_match:
                        ep = int(loss_match.group(1))
                        box_loss = float(loss_match.group(4))
                        self.current_training_epoch = ep
                        self.root.after(0, self._update_loss_chart, ep, box_loss)
                        self.root.after(0, self._append_excel_row, ep, box_loss, last_map_val)

                    map_match = map_regex.search(line)
                    if map_match:
                        last_map_val = float(map_match.group(1))
                        self.root.after(0, self._update_map_chart, self.current_training_epoch, last_map_val)

            proc.wait()
            rc = proc.returncode

            if rc == 0:
                self.root.after(0, self._training_complete)
            else:
                self.root.after(0, self._tlog, f"\nProcess exited with code {rc}")
                self.root.after(0, self._training_reset)
                self.root.after(0, self._set_status, "TRAINING FAILED", DANGER)

        except Exception as e:
            print(f"Training error: {e}")
            self.root.after(0, self._set_status, f"ERROR: {e}", DANGER)
            self.root.after(0, self._training_reset)

    _training_worker = _run_training_thread

    def _tlog(self, text):
        self.train_log.config(state="normal")
        self.train_log.insert("end", text + "\n")
        self.train_log.see("end")
        self.train_log.config(state="disabled")

    def _update_loss_chart(self, epoch, loss):
        if self.loss_epochs and epoch == self.loss_epochs[-1]:
            self.train_losses[-1] = loss
        else:
            self.loss_epochs.append(epoch)
            self.train_losses.append(loss)

        self.ax_loss.clear()
        self._style_axis(self.ax_loss, "Train Loss", "Epoch", "Loss")
        self.ax_loss.plot(self.loss_epochs, self.train_losses,
                          color=ACCENT, linewidth=1.5, marker="o", markersize=3)
        if len(self.train_losses) > 1:
            self.ax_loss.fill_between(self.loss_epochs, self.train_losses,
                                       alpha=0.08, color=ACCENT)
        self.chart_canvas.draw_idle()

    def _update_map_chart(self, epoch, mAP):
        if self.map_epochs and epoch == self.map_epochs[-1]:
            self.map_values[-1] = mAP
        else:
            self.map_epochs.append(epoch)
            self.map_values.append(mAP)

        self.ax_map.clear()
        self._style_axis(self.ax_map, "mAP@50", "Epoch", "mAP")
        self.ax_map.plot(self.map_epochs, self.map_values,
                         color=WARNING, linewidth=1.5, marker="o", markersize=3)
        if len(self.map_values) > 1:
            self.ax_map.fill_between(self.map_epochs, self.map_values,
                                      alpha=0.08, color=WARNING)
        self.chart_canvas.draw_idle()

    def _init_excel_log(self, save_dir, run_name):
        """Create a new Excel workbook for this training run."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Training Log"

            # Header row styling
            headers = ["Epoch", "Box Loss", "mAP@50", "Timestamp"]
            header_fill = PatternFill("solid", fgColor="0A0A0A")
            header_font = Font(name="Consolas", bold=True, color="00FF88", size=11)

            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
                ws.column_dimensions[get_column_letter(col)].width = 18

            # Metadata sheet
            ws_meta = wb.create_sheet("Run Info")
            ws_meta["A1"] = "SENTINEL AI — Training Run"
            ws_meta["A2"] = f"Run Name: {run_name}"
            ws_meta["A3"] = f"Started: {time.strftime('%Y-%m-%d %H:%M:%S')}"
            ws_meta["A4"] = f"Model: {self.model_var.get()}"
            ws_meta["A5"] = f"Epochs: {self.epochs_var.get()}"
            ws_meta["A6"] = f"Image Size: {self.imgsz_var.get()}"
            ws_meta["A7"] = f"Batch Size: {self.batch_var.get()}"

            os.makedirs(save_dir, exist_ok=True)
            self.excel_log_path = os.path.join(save_dir, f"{run_name}_training_log.xlsx")
            wb.save(self.excel_log_path)
            self._tlog(f"[EXCEL] Log started: {self.excel_log_path}")
        except Exception as e:
            self._tlog(f"[EXCEL] Could not create log: {e}")
            self.excel_log_path = None

    def _append_excel_row(self, epoch, loss, mAP):
        """Append one epoch row to the Excel log."""
        if not self.excel_log_path or not os.path.isfile(self.excel_log_path):
            return
        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Font

            wb = openpyxl.load_workbook(self.excel_log_path)
            ws = wb["Training Log"]
            row = ws.max_row + 1

            # Alternating row colors
            fill_color = "1A1A1A" if row % 2 == 0 else "111111"
            fill = PatternFill("solid", fgColor=fill_color)
            font = Font(name="Consolas", color="E0E0E0", size=10)

            values = [epoch, round(loss, 6), round(mAP, 6),
                      time.strftime('%H:%M:%S')]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.fill = fill
                cell.font = font

            wb.save(self.excel_log_path)
        except Exception as e:
            print(f"[EXCEL] Row append failed: {e}")

    def _training_complete(self):
        self._training_reset()
        self.btn_test.config(state="normal", bg=ACCENT, fg=BG)
        self._set_status("IDLE", ACCENT)
        self._export_training_graphs()
        self._tlog(
            "Training complete. "
            f"Model saved to {self.savedir_var.get()}; "
            f"Excel log: {os.path.basename(self.excel_log_path) if self.excel_log_path else 'N/A'}; "
            "graphs saved to models/ folder."
        )

    def _export_training_graphs(self):
        """Save loss and mAP charts as PNG files to the models/ folder."""
        try:
            import matplotlib.pyplot as plt

            graphs_dir = os.path.join(BASE_DIR, "models")
            os.makedirs(graphs_dir, exist_ok=True)
            ts = time.strftime("%Y%m%d_%H%M%S")

            # Loss graph
            if self.loss_epochs and self.train_losses:
                fig, ax = plt.subplots(figsize=(10, 5), facecolor="#0a0a0a")
                ax.set_facecolor("#1a1a1a")
                ax.plot(self.loss_epochs, self.train_losses,
                        color="#00ff88", linewidth=2, marker="o", markersize=4)
                ax.fill_between(self.loss_epochs, self.train_losses,
                                alpha=0.1, color="#00ff88")
                ax.set_title("SENTINEL AI — Training Loss", color="#00ff88",
                             fontsize=14, fontfamily="monospace")
                ax.set_xlabel("Epoch", color="#666666", fontfamily="monospace")
                ax.set_ylabel("Box Loss", color="#666666", fontfamily="monospace")
                ax.tick_params(colors="#666666")
                for spine in ax.spines.values():
                    spine.set_color("#666666")
                ax.grid(True, alpha=0.15, color="#666666")
                loss_path = os.path.join(graphs_dir, f"loss_curve_{ts}.png")
                fig.savefig(loss_path, dpi=150, bbox_inches="tight",
                            facecolor="#0a0a0a")
                plt.close(fig)
                self._tlog(f"[GRAPH] Loss curve saved: {loss_path}")

            # mAP graph
            if self.map_epochs and self.map_values:
                fig, ax = plt.subplots(figsize=(10, 5), facecolor="#0a0a0a")
                ax.set_facecolor("#1a1a1a")
                ax.plot(self.map_epochs, self.map_values,
                        color="#ffaa00", linewidth=2, marker="o", markersize=4)
                ax.fill_between(self.map_epochs, self.map_values,
                                alpha=0.1, color="#ffaa00")
                ax.set_title("SENTINEL AI — mAP@50", color="#ffaa00",
                             fontsize=14, fontfamily="monospace")
                ax.set_xlabel("Epoch", color="#666666", fontfamily="monospace")
                ax.set_ylabel("mAP@50", color="#666666", fontfamily="monospace")
                ax.tick_params(colors="#666666")
                for spine in ax.spines.values():
                    spine.set_color("#666666")
                ax.grid(True, alpha=0.15, color="#666666")
                map_path = os.path.join(graphs_dir, f"map_curve_{ts}.png")
                fig.savefig(map_path, dpi=150, bbox_inches="tight",
                            facecolor="#0a0a0a")
                plt.close(fig)
                self._tlog(f"[GRAPH] mAP curve saved: {map_path}")

        except Exception as e:
            self._tlog(f"[GRAPH] Export failed: {e}")

    def _training_reset(self):
        self.btn_train.config(state="normal")
        self.btn_stop_train.config(state="disabled")
        self.training_process = None

    def _stop_training(self):
        if self.training_process:
            try:
                self.training_process.terminate()
            except Exception:
                pass
            self._set_status("TRAINING STOPPED", DANGER)
            self._training_reset()

    def _test_model(self):
        filepath = filedialog.askopenfilename(
            title="Select Test Image",
            filetypes=[("Images", "*.jpg *.jpeg *.png *.bmp")]
        )
        if not filepath:
            return

        model, status = load_person_model(BASE_DIR)
        if model is None:
            self._set_status(f"MODEL ERROR: {status}", DANGER)
            self._tlog(f"Model test unavailable: {status}")
            return

        results = model(filepath, verbose=False)
        result_img = results[0].plot()

        win = tk.Toplevel(self.root)
        win.title("Inference Result")
        win.configure(bg=BG)

        rgb = cv2.cvtColor(result_img, cv2.COLOR_BGR2RGB)
        h, w = rgb.shape[:2]
        scale = min(900 / w, 650 / h, 1.0)
        resized = cv2.resize(rgb, (int(w * scale), int(h * scale)))
        photo = ImageTk.PhotoImage(Image.fromarray(resized))

        lbl = tk.Label(win, image=photo, bg=BG)
        lbl.image = photo
        lbl.pack(padx=10, pady=10)

    # ═════════════════════════════════════════════════════════════════════════
    # TAB 3 — VIDEO DEMO
    # ═════════════════════════════════════════════════════════════════════════

    def _build_video_tab(self):
        # ── Controls bar ──
        ctrl = tk.Frame(self.tab_video, bg=PANEL, height=55)
        ctrl.pack(fill="x", padx=10, pady=(10, 5))
        ctrl.pack_propagate(False)

        tk.Button(ctrl, text="\U0001F4C1 UPLOAD VIDEO", font=FONT_BODY,
                  bg=CARD, fg=TEXT, relief="flat", cursor="hand2",
                  activebackground=PANEL, command=self._upload_video
                  ).pack(side="left", padx=(10, 5), pady=8)

        self.video_fname = tk.Label(ctrl, text="No file selected", font=FONT_SMALL,
                                     fg=MUTED, bg=PANEL)
        self.video_fname.pack(side="left", padx=5)
        self.video_time_label = tk.Label(ctrl, text="00:00 / 00:00", font=FONT_SMALL,
                                         fg=MUTED, bg=PANEL)
        self.video_time_label.pack(side="left", padx=(8, 5))

        self.btn_vid_play = tk.Button(ctrl, text="\u25B6 PLAY", font=FONT_BODY,
                                       bg=ACCENT, fg=BG, relief="flat", cursor="hand2",
                                       activebackground="#00cc6a",
                                       command=self._play_video, state="disabled")
        self.btn_vid_play.pack(side="left", padx=5, pady=8)

        self.btn_vid_stop = tk.Button(ctrl, text="\u25A0 STOP", font=FONT_BODY,
                                       bg=DANGER, fg=TEXT, relief="flat", cursor="hand2",
                                       activebackground="#cc2929",
                                       command=self._stop_video, state="disabled")
        self.btn_vid_stop.pack(side="left", padx=5, pady=8)

        tk.Label(ctrl, text="CONFIDENCE", font=FONT_SMALL, fg=MUTED,
                 bg=PANEL).pack(side="left", padx=(20, 2))
        self.vid_conf = tk.DoubleVar(value=0.25)
        tk.Scale(ctrl, from_=0.1, to=0.9, resolution=0.05, orient="horizontal",
                 variable=self.vid_conf, bg=PANEL, fg=ACCENT, troughcolor=CARD,
                 highlightthickness=0, font=FONT_SMALL, length=110, sliderlength=14
                 ).pack(side="left", padx=2)

        # ── Main Layout (Split Frame) ──
        main_layout = tk.Frame(self.tab_video, bg=BG)
        main_layout.pack(fill="both", expand=True, padx=10, pady=5)

        # Left Column (Video canvas + charts + alert log)
        left_col = tk.Frame(main_layout, bg=BG)
        left_col.pack(side="left", fill="both", expand=True)

        self.vid_canvas = tk.Canvas(left_col, bg="black", highlightthickness=0)
        self.vid_canvas.pack(pady=5, fill="both", expand=True)
        self.video_progress = ttk.Progressbar(left_col, mode="determinate",
                                              style="Horizontal.TProgressbar")
        self.video_progress.pack(fill="x", pady=(0, 6))

        hist_frame = tk.Frame(left_col, bg=CARD, height=105)
        hist_frame.pack(pady=(0, 5), fill="x")
        hist_frame.pack_propagate(False)

        tk.Label(hist_frame, text="CONFIDENCE DISTRIBUTION", font=FONT_SMALL,
                 fg=MUTED, bg=CARD).pack(anchor="w", padx=8, pady=(4, 0))

        self.vid_hist_fig = Figure(figsize=(8, 0.75), dpi=100, facecolor=CARD)
        self.vid_hist_fig.subplots_adjust(left=0.05, right=0.98, top=0.85, bottom=0.25)
        self.vid_hist_ax = self.vid_hist_fig.add_subplot(111)
        self.vid_hist_ax.set_facecolor(BG)
        self.vid_hist_ax.tick_params(colors=MUTED, labelsize=7)
        for spine in self.vid_hist_ax.spines.values():
            spine.set_color(MUTED)
            spine.set_linewidth(0.4)

        self.vid_hist_canvas = FigureCanvasTkAgg(self.vid_hist_fig, master=hist_frame)
        self.vid_hist_canvas.get_tk_widget().pack(fill="both", expand=True)
        self.vid_hist_canvas.draw()
        self.vid_conf_history = []   # rolling confidence values

        tk.Label(left_col, text="ALERT LOG", font=FONT_SMALL, fg=MUTED,
                 bg=BG).pack(anchor="w")
        self.vid_alert_log = tk.Text(left_col, bg=CARD, fg=DANGER, font=FONT_SMALL,
                                      relief="flat", height=5, padx=10, pady=5,
                                      highlightthickness=0)
        self.vid_alert_log.pack(fill="x")
        self.vid_alert_log.config(state="disabled")

        # Right Column (Cyber-HUD Sidebar Behavior Panel)
        right_col = tk.Frame(main_layout, bg=PANEL, width=320)
        right_col.pack(side="right", fill="both", expand=False, padx=(10, 0))
        right_col.pack_propagate(False)

        tk.Label(right_col, text="ANALYSIS DETAILS", font=FONT_HEADING,
                 fg=ACCENT, bg=PANEL).pack(pady=(15, 10))

        # Active Targets Display
        tk.Label(right_col, text="ACTIVE TARGETS", font=FONT_SMALL, fg=MUTED,
                 bg=PANEL).pack(anchor="w", padx=15, pady=(10, 2))
        self.vid_targets_log = tk.Text(right_col, bg=CARD, fg=TEXT, font=FONT_BODY,
                                       relief="flat", height=10, padx=10, pady=5,
                                       highlightthickness=0)
        self.vid_targets_log.pack(fill="x", padx=10, pady=2)
        self.vid_targets_log.config(state="disabled")

        # Global Risk Meter
        tk.Label(right_col, text="GLOBAL RISK SEVERITY", font=FONT_SMALL, fg=MUTED,
                 bg=PANEL).pack(anchor="w", padx=15, pady=(10, 2))
        self.vid_risk_lbl = tk.Label(right_col, text="NORMAL", font=("Consolas", 20, "bold"),
                                     fg=ACCENT, bg=CARD, pady=12)
        self.vid_risk_lbl.pack(fill="x", padx=10, pady=2)

        # Recent Behavior Events
        tk.Label(right_col, text="RECENT BEHAVIOR EVENTS", font=FONT_SMALL, fg=MUTED,
                 bg=PANEL).pack(anchor="w", padx=15, pady=(10, 2))
        self.vid_events_log = tk.Text(right_col, bg=CARD, fg=WARNING, font=FONT_SMALL,
                                      relief="flat", height=6, padx=10, pady=5,
                                      highlightthickness=0)
        self.vid_events_log.pack(fill="both", expand=True, padx=10, pady=(2, 15))
        self.vid_events_log.config(state="disabled")

        self.video_path = None
        self.vid_tracked_entities = {}
        self.vid_fcount = 0

    def _upload_video(self):
        path = filedialog.askopenfilename(
            title="Select Video File",
            filetypes=[("Video Files", "*.mp4 *.avi *.mov *.mpg"), ("All Files", "*.*")]
        )
        if path:
            self.video_path = path
            self.video_fname.config(text=os.path.basename(path), fg=TEXT)
            self.video_source_status = f"Selected {os.path.basename(path)}"
            self.btn_vid_play.config(state="normal")

    def _play_video(self):
        if not self.video_path:
            return

        person_model, status = load_person_model(BASE_DIR)
        pose_model = load_pose_model(BASE_DIR)
        self.model_status_text = status
        self.pose_status_text = "Available" if pose_model is not None else "Unavailable"
        if person_model is None:
            self._canvas_msg(self.vid_canvas, f"Models failed to load: {status}")
            return

        self.video_running = True
        self.vid_tracked_entities = {}
        self.vid_fcount = 0
        self.btn_vid_play.config(state="disabled")
        self.btn_vid_stop.config(state="normal")
        self.video_source_status = f"Processing {os.path.basename(self.video_path)}"
        self._set_status("DETECTING", WARNING)

        threading.Thread(target=self._video_worker, args=(person_model, pose_model), daemon=True).start()

    def _video_worker(self, person_model, pose_model):
        cap = None
        try:
            if not self.video_path or not os.path.isfile(self.video_path):
                self.root.after(0, self._canvas_msg, self.vid_canvas, "Video file not found or invalid.")
                self.root.after(0, self._set_status, "FILE NOT FOUND", DANGER)
                return

            cap = cv2.VideoCapture(self.video_path)
            if not cap or not cap.isOpened():
                self.root.after(0, self._canvas_msg, self.vid_canvas, f"Cannot open video:\n{os.path.basename(self.video_path)}")
                self.root.after(0, self._set_status, "CANNOT OPEN VIDEO", DANGER)
                return

            fps = cap.get(cv2.CAP_PROP_FPS) or 30
            frame_count = cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0
            self.video_duration = (frame_count / fps) if fps else 0.0
            delay = 1.0 / max(fps, 1)
            perf_last = time.perf_counter()

            while self.video_running and cap.isOpened():
                ret, frame = cap.read()
                if not ret or frame is None or frame.size == 0:
                    break

                self.vid_fcount += 1
                video_time = cap.get(cv2.CAP_PROP_POS_MSEC) / 1000.0
                if video_time <= 0:
                    video_time = self.vid_fcount / max(fps, 1)
                self.video_position = video_time
                perf_now = time.perf_counter()
                dt = perf_now - perf_last
                perf_last = perf_now
                if dt > 0:
                    self.vid_processing_fps = 1.0 / dt
                
                person_dets = detect_persons(frame, person_model, self.vid_conf.get())
                merged_dets = iou_deduplicate(person_dets)

                # Tracking match
                self.vid_tracked_entities = match_entities(
                    merged_dets, self.vid_tracked_entities, current_time=video_time
                )
                
                # Human posture estimation & temporal fall check
                fall_alerts = estimate_postures(
                    frame, self.vid_tracked_entities, pose_model, current_time=video_time
                )

                # Behavior analysis and overlay drawing (renders skeleton and telemetry)
                annotated, new_alerts = draw_overlays(
                    frame, self.vid_tracked_entities, self.vid_fcount, current_time=video_time
                )
                new_alerts = fall_alerts + new_alerts
                self._last_frame = annotated.copy()

                # Update live HUD statistics on every frame
                self.root.after(0, self._update_hud_sidebar, "vid", self.vid_tracked_entities)

                self.root.after(0, self._update_conf_histogram,
                                person_dets,
                                self.vid_hist_ax, self.vid_hist_canvas,
                                self.vid_conf_history)

                # Process newly triggered suspicious behavior alerts
                for alert in new_alerts:
                    is_high_threat = (alert["severity"] in ("HIGH", "CRITICAL"))
                    
                    if is_high_threat or alert["behavior"] == "POSSIBLE_FALL":
                        snapshot = annotated.copy()
                        alert["snapshot_path"] = self._save_incident_snapshot(
                            snapshot, alert["behavior"], alert["track_id"], alert["severity"])
                        self.root.after(0, self._play_alert_sound)

                    self.root.after(0, self._fire_behavior_alert, "vid", alert)

                self.root.after(0, self._update_video_progress)
                self.root.after(0, self._show_frame, annotated, self.vid_canvas)
                time.sleep(delay)

        except Exception as e:
            print(f"Video error: {e}")
            self.root.after(0, self._set_status, f"ERROR: {e}", DANGER)
        finally:
            if cap is not None:
                try:
                    cap.release()
                except Exception:
                    pass
            self.root.after(0, self._video_stopped)

    def _stop_video(self):
        self.video_running = False
        self._stop_alert_sound()

    def _video_stopped(self):
        self.video_running = False
        self.video_source_status = "Stopped"
        self.btn_vid_play.config(state="normal")
        self.btn_vid_stop.config(state="disabled")
        self._set_status("IDLE", ACCENT)

    def _update_video_progress(self):
        if hasattr(self, "video_time_label"):
            self.video_time_label.config(
                text=f"{self._format_seconds(self.video_position)} / {self._format_seconds(self.video_duration)}"
            )
        if hasattr(self, "video_progress"):
            if self.video_duration > 0:
                self.video_progress.config(maximum=self.video_duration, value=min(self.video_position, self.video_duration))
            else:
                self.video_progress.config(value=0)

    # ═════════════════════════════════════════════════════════════════════════
    # TAB 4 — LIVE CAMERA
    # ═════════════════════════════════════════════════════════════════════════

    def _build_camera_tab(self):
        # ── Controls bar ──
        ctrl = tk.Frame(self.tab_camera, bg=PANEL, height=55)
        ctrl.pack(fill="x", padx=10, pady=(10, 5))
        ctrl.pack_propagate(False)

        self.btn_cam_start = tk.Button(ctrl, text="\u25B6 START CAMERA", font=FONT_BODY,
                                        bg=ACCENT, fg=BG, relief="flat", cursor="hand2",
                                        activebackground="#00cc6a",
                                        command=self._start_camera)
        self.btn_cam_start.pack(side="left", padx=(10, 5), pady=8)

        self.btn_cam_stop = tk.Button(ctrl, text="\u25A0 STOP CAMERA", font=FONT_BODY,
                                       bg=DANGER, fg=TEXT, relief="flat", cursor="hand2",
                                       activebackground="#cc2929",
                                       command=self._stop_camera, state="disabled")
        self.btn_cam_stop.pack(side="left", padx=5, pady=8)
        self.cam_status_inline = tk.Label(ctrl, text="Disconnected", font=FONT_SMALL,
                                          fg=MUTED, bg=PANEL)
        self.cam_status_inline.pack(side="left", padx=(10, 4))
        self.cam_fps_inline = tk.Label(ctrl, text="0.0 FPS", font=FONT_SMALL,
                                       fg=MUTED, bg=PANEL)
        self.cam_fps_inline.pack(side="left", padx=4)

        tk.Label(ctrl, text="CAMERA", font=FONT_SMALL, fg=MUTED,
                 bg=PANEL).pack(side="left", padx=(20, 5))
        self.cam_idx_var = tk.StringVar(value="0")
        ttk.Combobox(ctrl, textvariable=self.cam_idx_var,
                      values=["0", "1", "2"], state="readonly",
                      font=FONT_BODY, width=3).pack(side="left", padx=2)

        tk.Label(ctrl, text="CONFIDENCE", font=FONT_SMALL, fg=MUTED,
                 bg=PANEL).pack(side="left", padx=(20, 2))
        self.cam_conf = tk.DoubleVar(value=0.25)
        tk.Scale(ctrl, from_=0.1, to=0.9, resolution=0.05, orient="horizontal",
                 variable=self.cam_conf, bg=PANEL, fg=ACCENT, troughcolor=CARD,
                 highlightthickness=0, font=FONT_SMALL, length=110, sliderlength=14
                 ).pack(side="left", padx=2)

        # ── Main Layout (Split Frame) ──
        main_layout = tk.Frame(self.tab_camera, bg=BG)
        main_layout.pack(fill="both", expand=True, padx=10, pady=5)

        # Left Column
        left_col = tk.Frame(main_layout, bg=BG)
        left_col.pack(side="left", fill="both", expand=True)

        self.cam_canvas = tk.Canvas(left_col, bg="black", highlightthickness=0)
        self.cam_canvas.pack(pady=5, fill="both", expand=True)

        hist_frame = tk.Frame(left_col, bg=CARD, height=105)
        hist_frame.pack(pady=(0, 5), fill="x")
        hist_frame.pack_propagate(False)

        tk.Label(hist_frame, text="CONFIDENCE DISTRIBUTION", font=FONT_SMALL,
                 fg=MUTED, bg=CARD).pack(anchor="w", padx=8, pady=(4, 0))

        self.cam_hist_fig = Figure(figsize=(8, 0.75), dpi=100, facecolor=CARD)
        self.cam_hist_fig.subplots_adjust(left=0.05, right=0.98, top=0.85, bottom=0.25)
        self.cam_hist_ax = self.cam_hist_fig.add_subplot(111)
        self.cam_hist_ax.set_facecolor(BG)
        self.cam_hist_ax.tick_params(colors=MUTED, labelsize=7)
        for spine in self.cam_hist_ax.spines.values():
            spine.set_color(MUTED)
            spine.set_linewidth(0.4)

        self.cam_hist_canvas = FigureCanvasTkAgg(self.cam_hist_fig, master=hist_frame)
        self.cam_hist_canvas.get_tk_widget().pack(fill="both", expand=True)
        self.cam_hist_canvas.draw()
        self.cam_conf_history = []

        tk.Label(left_col, text="ALERT LOG", font=FONT_SMALL, fg=MUTED,
                 bg=BG).pack(anchor="w")
        self.cam_alert_log = tk.Text(left_col, bg=CARD, fg=DANGER, font=FONT_SMALL,
                                      relief="flat", height=5, padx=10, pady=5,
                                      highlightthickness=0)
        self.cam_alert_log.pack(fill="x")
        self.cam_alert_log.config(state="disabled")

        # Right Column (Cyber-HUD Sidebar Behavior Panel)
        right_col = tk.Frame(main_layout, bg=PANEL, width=320)
        right_col.pack(side="right", fill="both", expand=False, padx=(10, 0))
        right_col.pack_propagate(False)

        tk.Label(right_col, text="TARGET DETAILS", font=FONT_HEADING,
                 fg=ACCENT, bg=PANEL).pack(pady=(15, 10))

        tk.Label(right_col, text="ACTIVE TARGETS", font=FONT_SMALL, fg=MUTED,
                 bg=PANEL).pack(anchor="w", padx=15, pady=(10, 2))
        self.cam_targets_log = tk.Text(right_col, bg=CARD, fg=TEXT, font=FONT_BODY,
                                       relief="flat", height=10, padx=10, pady=5,
                                       highlightthickness=0)
        self.cam_targets_log.pack(fill="x", padx=10, pady=2)
        self.cam_targets_log.config(state="disabled")

        tk.Label(right_col, text="GLOBAL RISK SEVERITY", font=FONT_SMALL, fg=MUTED,
                 bg=PANEL).pack(anchor="w", padx=15, pady=(10, 2))
        self.cam_risk_lbl = tk.Label(right_col, text="NORMAL", font=("Consolas", 20, "bold"),
                                     fg=ACCENT, bg=CARD, pady=12)
        self.cam_risk_lbl.pack(fill="x", padx=10, pady=2)

        tk.Label(right_col, text="RECENT BEHAVIOR EVENTS", font=FONT_SMALL, fg=MUTED,
                 bg=PANEL).pack(anchor="w", padx=15, pady=(10, 2))
        self.cam_events_log = tk.Text(right_col, bg=CARD, fg=WARNING, font=FONT_SMALL,
                                      relief="flat", height=6, padx=10, pady=5,
                                      highlightthickness=0)
        self.cam_events_log.pack(fill="both", expand=True, padx=10, pady=(2, 15))
        self.cam_events_log.config(state="disabled")

        self.cam_tracked_entities = {}
        self.cam_fcount = 0

    def _start_camera(self):
        person_model, status = load_person_model(BASE_DIR)
        pose_model = load_pose_model(BASE_DIR)
        self.model_status_text = status
        self.pose_status_text = "Available" if pose_model is not None else "Unavailable"
        if person_model is None:
            self.root.after(50, lambda: self._canvas_msg(
                self.cam_canvas, f"Models failed to load: {status}"))
            return

        cam_idx_str = self.cam_idx_var.get()
        try:
            cam_idx = int(cam_idx_str)
        except (ValueError, TypeError):
            cam_idx = 0

        self.camera_running = True
        self.cam_tracked_entities = {}
        self.cam_fcount = 0
        self.btn_cam_start.config(state="disabled")
        self.btn_cam_stop.config(state="normal")
        self.camera_source_status = f"Connecting camera {cam_idx}"
        self.cam_status_inline.config(text=self.camera_source_status, fg=WARNING)
        self._set_status("CONNECTING TO CAMERA...", WARNING)

        threading.Thread(target=self._camera_worker,
                         args=(person_model, pose_model, cam_idx), daemon=True).start()

    def _camera_worker(self, person_model, pose_model, cam_idx):
        cap = None
        try:
            cap = cv2.VideoCapture(cam_idx)
            if not cap or not cap.isOpened():
                self.root.after(0, self._canvas_msg, self.cam_canvas,
                                f"\u26A0 No camera detected at index {cam_idx}\nPlease check webcam connection or select another index.")
                self.root.after(0, self._set_status, f"CAMERA NOT FOUND (idx {cam_idx})", DANGER)
                self.camera_source_status = f"Camera {cam_idx} unavailable"
                return

            # Test-read first frame
            ret, test_frame = cap.read()
            if not ret or test_frame is None:
                self.root.after(0, self._canvas_msg, self.cam_canvas,
                                f"\u26A0 Failed to read feed from camera index {cam_idx}")
                self.root.after(0, self._set_status, f"CAMERA READ FAILED (idx {cam_idx})", DANGER)
                self.camera_source_status = f"Camera {cam_idx} read failed"
                return

            self.root.after(0, self._set_status, "CAMERA ACTIVE", ACCENT)
            self.camera_source_status = f"Camera {cam_idx} active"
            self.root.after(0, self._update_camera_inline_status)

            failed_reads = 0
            perf_last = time.perf_counter()
            while self.camera_running and cap.isOpened():
                ret, frame = cap.read()
                if not ret or frame is None or frame.size == 0:
                    failed_reads += 1
                    if failed_reads > 30:
                        self.root.after(0, self._canvas_msg, self.cam_canvas,
                                        f"\u26A0 Camera feed lost or disconnected (index {cam_idx})")
                        self.root.after(0, self._set_status, "CAMERA DISCONNECTED", DANGER)
                        self.camera_source_status = f"Camera {cam_idx} disconnected"
                        self.root.after(0, self._update_camera_inline_status)
                        break
                    time.sleep(0.03)
                    continue

                failed_reads = 0
                self.cam_fcount += 1
                frame_time = time.time()
                perf_now = time.perf_counter()
                dt = perf_now - perf_last
                perf_last = perf_now
                if dt > 0:
                    self.cam_processing_fps = 1.0 / dt
                
                person_dets = detect_persons(frame, person_model, self.cam_conf.get())
                merged_dets = iou_deduplicate(person_dets)

                self.cam_tracked_entities = match_entities(
                    merged_dets, self.cam_tracked_entities, current_time=frame_time
                )
                
                # Human posture estimation & temporal fall check
                fall_alerts = estimate_postures(
                    frame, self.cam_tracked_entities, pose_model, current_time=frame_time
                )

                # Behavior analysis and overlay drawing (renders skeleton and telemetry)
                annotated, new_alerts = draw_overlays(
                    frame, self.cam_tracked_entities, self.cam_fcount, current_time=frame_time
                )
                new_alerts = fall_alerts + new_alerts
                self._last_frame = annotated.copy()

                self.root.after(0, self._update_hud_sidebar, "cam", self.cam_tracked_entities)

                self.root.after(0, self._update_conf_histogram,
                                person_dets,
                                self.cam_hist_ax, self.cam_hist_canvas,
                                self.cam_conf_history)

                for alert in new_alerts:
                    is_high_threat = (alert["severity"] in ("HIGH", "CRITICAL"))
                    
                    if is_high_threat or alert["behavior"] == "POSSIBLE_FALL":
                        snapshot = annotated.copy()
                        alert["snapshot_path"] = self._save_incident_snapshot(
                            snapshot, alert["behavior"], alert["track_id"], alert["severity"])
                        self.root.after(0, self._play_alert_sound)

                    self.root.after(0, self._fire_behavior_alert, "cam", alert)

                self.root.after(0, self._update_camera_inline_status)
                self.root.after(0, self._show_frame, annotated, self.cam_canvas)
                time.sleep(0.03)  # ~30 fps cap

        except Exception as e:
            print(f"Camera error: {e}")
            self.root.after(0, self._set_status, f"ERROR: {e}", DANGER)
        finally:
            if cap is not None:
                try:
                    cap.release()
                except Exception:
                    pass
            self.root.after(0, self._camera_stopped)

    def _stop_camera(self):
        self.camera_running = False
        self._stop_alert_sound()

    def _camera_stopped(self):
        self.camera_running = False
        if "active" in self.camera_source_status.lower():
            self.camera_source_status = "Disconnected"
        self._update_camera_inline_status()
        self.btn_cam_start.config(state="normal")
        self.btn_cam_stop.config(state="disabled")
        self._set_status("IDLE", ACCENT)

    def _update_camera_inline_status(self):
        if hasattr(self, "cam_status_inline"):
            color = ACCENT if self.camera_running and "active" in self.camera_source_status.lower() else (
                DANGER if "failed" in self.camera_source_status.lower() or "unavailable" in self.camera_source_status.lower()
                or "disconnected" in self.camera_source_status.lower() else WARNING
            )
            self.cam_status_inline.config(text=self.camera_source_status, fg=color)
        if hasattr(self, "cam_fps_inline"):
            self.cam_fps_inline.config(text=f"{self.cam_processing_fps:.1f} FPS")

    # ═════════════════════════════════════════════════════════════════════════
    # SHARED DISPLAY HELPERS
    # ═════════════════════════════════════════════════════════════════════════

    def _show_frame(self, frame, canvas):
        """Render a cv2 frame onto a Tk canvas, centred and aspect-ratio correct."""
        try:
            canvas.update_idletasks()
            cw = canvas.winfo_width()
            ch = canvas.winfo_height()
            if cw < 10 or ch < 10:
                return

            photo = frame_to_tk(frame, cw, ch)
            if photo is None:
                return
            canvas.delete("all")
            canvas.create_image(cw // 2, ch // 2, image=photo, anchor="center")
            canvas._photo = photo  # prevent GC
        except (tk.TclError, Exception):
            pass

    def _canvas_msg(self, canvas, text):
        """Show a centred text message on a canvas."""
        try:
            canvas.update_idletasks()
            cw = canvas.winfo_width()
            ch = canvas.winfo_height()
            canvas.delete("all")
            canvas.create_text(max(cw // 2, 100), max(ch // 2, 50),
                               text=text, fill=DANGER, font=FONT_HEADING,
                               width=max(cw - 40, 200))
        except (tk.TclError, Exception):
            pass

    def _update_conf_histogram(self, person_dets, hist_ax, hist_canvas, conf_history):
        """Update the confidence histogram with latest detections."""
        try:
            new_confs = [d[4] for d in person_dets]
            conf_history.extend(new_confs)
            if len(conf_history) > 150:
                del conf_history[:-150]

            hist_ax.clear()
            hist_ax.set_facecolor(BG)
            for spine in hist_ax.spines.values():
                spine.set_color(MUTED)
                spine.set_linewidth(0.4)
            hist_ax.tick_params(colors=MUTED, labelsize=7)

            if conf_history:
                bins = [i/10 for i in range(1, 11)]
                counts = [0] * 9
                for c in conf_history:
                    idx = min(int((c - 0.1) / 0.1), 8)
                    if 0 <= idx < 9:
                        counts[idx] += 1
                bar_colors = []
                for i, b in enumerate(bins[:-1]):
                    mid = b + 0.05
                    if mid < 0.5:
                        bar_colors.append("#ffaa00")
                    elif mid < 0.75:
                        bar_colors.append("#00aaff")
                    else:
                        bar_colors.append("#00ff88")
                hist_ax.bar([b + 0.05 for b in bins[:-1]], counts,
                            width=0.08, color=bar_colors, alpha=0.85)
                hist_ax.set_xlim(0.1, 1.0)
                hist_ax.set_xlabel("Confidence", color=MUTED,
                                   fontsize=7, fontfamily="monospace")

            hist_canvas.draw_idle()
        except Exception as e:
            print(f"Histogram error: {e}")

    def _update_hud_sidebar(self, stream_type, tracked_entities):
        """Update the active target list and global threat severity gauge in the right Cyber-HUD sidebar panel."""
        targets_text = ""
        max_risk_score = 0
        global_severity = "NORMAL"
        severity_color = ACCENT

        # Retrieve text widget and risk label depending on tab stream type
        targets_log = getattr(self, f"{stream_type}_targets_log", None)
        risk_lbl = getattr(self, f"{stream_type}_risk_lbl", None)

        if not targets_log or not risk_lbl:
            return

        # Build lines for each tracked entity
        for key, info in sorted(tracked_entities.items()):
            if info.get("alert_dismissed", False):
                continue
            
            cls_name = info["class_name"]
            track_id = info["track_id"]
            threats = info.get("active_threats", set())
            risk_score, severity, color = calculate_entity_risk(info)

            if is_person(cls_name):
                dwell = int(info.get("dwell_time", 0))
                meander = float(info.get("meander_ratio", 1.0))
                revisits = int(info.get("revisit_count", 0))
                posture = info.get("posture", "UNKNOWN")
                pose_conf = int(info.get("pose_confidence", 0.0) * 100)
                b_state = str(info.get("behavior_state", "NORMAL"))
                behaviour_label = {
                    "ALERT": "SUSPICIOUS",
                    "CRITICAL": "SUSPICIOUS",
                    "HIGH": "SUSPICIOUS",
                    "COOLING_DOWN": "COOLING",
                }.get(b_state, b_state)
                targets_text += (
                    f"PERSON #{track_id} | Posture: {posture} ({pose_conf}%)\n"
                    f"  Dwell: {dwell}s | Meander: {meander:.1f} | Revisits: {revisits}\n"
                    f"  Behaviour: {behaviour_label} | Risk: {severity}\n"
                )
            else:
                continue
            
            # Keep track of global maximum risk severity
            if risk_score > max_risk_score:
                max_risk_score = risk_score
                global_severity = severity
                severity_color = color

        # Update HUD active targets panel
        targets_log.config(state="normal")
        targets_log.delete("1.0", "end")
        targets_log.insert("1.0", targets_text)
        targets_log.config(state="disabled")

        # Convert severity color to Hex for Tkinter safely
        if isinstance(severity_color, str) and severity_color.startswith("#"):
            hex_color = severity_color
        elif isinstance(severity_color, (tuple, list)) and len(severity_color) == 3:
            b, g, r = severity_color
            hex_color = f"#{int(r):02x}{int(g):02x}{int(b):02x}"
        else:
            hex_color = ACCENT

        # Update HUD risk gauge panel
        risk_lbl.config(text=f"{global_severity} ({max_risk_score}%)", fg=hex_color)

    def _format_seconds(self, seconds):
        seconds = max(0, int(seconds or 0))
        return f"{seconds // 60:02d}:{seconds % 60:02d}"

    def _severity_color(self, severity):
        return {
            "LOW": ACCENT,
            "MEDIUM": WARNING,
            "HIGH": WARNING,
            "CRITICAL": DANGER,
        }.get(str(severity).upper(), TEXT)

    def _incident_summary(self, alert):
        behavior = alert.get("behavior", "EVENT")
        evidence = alert.get("evidence") or {}
        if behavior == "POSSIBLE_FALL":
            return (
                f"Possible fall, not confirmed fall. Posture={evidence.get('posture', 'LYING_DOWN')}, "
                f"pose confidence={float(evidence.get('pose_confidence', 0.0)):.2f}"
            )
        if behavior in ("LOITERING", "SUSPICIOUS_LOITERING", "LONG_DWELL", "REPEATED_MOVEMENT"):
            return (
                "Behavioural risk evidence: "
                f"dwell={evidence.get('dwell_time', '?')}s, "
                f"meander={evidence.get('meander_ratio', '?')}, "
                f"revisits={evidence.get('revisit_count', '?')}. "
                "This is a risk signal, not a criminality probability."
            )
        return str(alert.get("description") or alert.get("event_name") or behavior)

    def _record_incident(self, stream_type, alert, alert_msg):
        category = alert.get("behavior", "EVENT")
        severity = alert.get("severity", "LOW")
        track_id = alert.get("track_id", "?")
        class_name = alert.get("class_name", "object")
        incident = {
            "id": len(self.incidents) + 1,
            "time": time.strftime("%H:%M:%S"),
            "date": time.strftime("%Y-%m-%d"),
            "source": "Live" if stream_type == "cam" else "Video",
            "category": category,
            "severity": severity,
            "track_id": track_id,
            "class_name": class_name,
            "object_label": f"{class_name.upper()}#{track_id}",
            "summary": self._incident_summary(alert),
            "message": alert_msg.strip(),
            "snapshot_path": alert.get("snapshot_path"),
            "acknowledged": False,
            "dismissed": False,
        }
        self.incidents.append(incident)
        self.unread_incidents += 1
        self._update_incident_badges()
        self._refresh_incidents_view()
        return incident

    def _update_incident_badges(self):
        if hasattr(self, "nav_incident_badge"):
            suffix = "incident" if self.unread_incidents == 1 else "incidents"
            self.nav_incident_badge.config(
                text=f"{self.unread_incidents} new {suffix}",
                fg=DANGER if self.unread_incidents else MUTED
            )

    def _refresh_incidents_view(self):
        if not hasattr(self, "incident_tree"):
            return
        category_filter = self.incident_category_filter.get()
        severity_filter = self.incident_severity_filter.get()

        categories = ["All"] + sorted({i["category"] for i in self.incidents})
        try:
            self.incident_category_combo.config(values=categories)
        except Exception:
            pass

        self.incident_tree.delete(*self.incident_tree.get_children())
        for incident in reversed(self.incidents):
            if category_filter != "All" and incident["category"] != category_filter:
                continue
            if severity_filter != "All" and incident["severity"] != severity_filter:
                continue
            status = "Dismissed" if incident["dismissed"] else ("Acknowledged" if incident["acknowledged"] else "New")
            self.incident_tree.insert(
                "", "end", iid=str(incident["id"]),
                values=(
                    incident["time"], incident["source"], incident["category"],
                    incident["severity"], incident["object_label"], status, incident["summary"]
                )
            )
        self._show_selected_incident_details()

    def _selected_incident(self):
        if not hasattr(self, "incident_tree"):
            return None
        selected = self.incident_tree.selection()
        if not selected:
            return None
        inc_id = int(selected[0])
        return next((i for i in self.incidents if i["id"] == inc_id), None)

    def _show_selected_incident_details(self):
        if not hasattr(self, "incident_details"):
            return
        incident = self._selected_incident()
        if not incident:
            self._set_text(self.incident_details, "Select an incident to inspect details.")
            return
        details = (
            f"{incident['date']} {incident['time']} | {incident['source']} | {incident['severity']}\n"
            f"Target: {incident['object_label']} | Category: {incident['category']}\n"
            f"Evidence: {incident['summary']}\n"
            f"Snapshot: {incident.get('snapshot_path') or 'Not saved for this event'}"
        )
        self._set_text(self.incident_details, details)

    def _ack_selected_incident(self):
        incident = self._selected_incident()
        if not incident:
            return
        incident["acknowledged"] = True
        self._stop_alert_sound()
        self._refresh_incidents_view()

    def _dismiss_selected_incident(self):
        incident = self._selected_incident()
        if not incident:
            return
        incident["dismissed"] = True
        incident["acknowledged"] = True
        self._dismiss_target_alert(incident["class_name"], incident["track_id"])
        self._stop_alert_sound()
        self._refresh_incidents_view()

    def _dismiss_target_alert(self, alert_type, target_id):
        for ents in [getattr(self, "vid_tracked_entities", {}),
                     getattr(self, "cam_tracked_entities", {})]:
            for _key, info in ents.items():
                if info.get("class_name") == alert_type and info.get("track_id") == target_id:
                    info["alert_dismissed"] = True
                    info.setdefault("active_threats", set()).clear()

    def _fire_behavior_alert(self, stream_type, alert):
        """
        Log alerts in dashboard widgets, update statistics counters, and record a
        non-modal incident entry for review and acknowledgement.
        """
        ts = time.strftime("%H:%M:%S")
        behavior = alert["behavior"]
        severity = alert["severity"]
        track_id = alert["track_id"]
        class_name = alert.get("class_name", "object")
        event_name = alert.get("event_name", behavior)
        evidence = alert.get("evidence") or {}
        evidence_note = ""
        if evidence:
            evidence_note = (
                f" [Dwell={evidence.get('dwell_time', '?')}s, "
                f"Meander={evidence.get('meander_ratio', '?')}, "
                f"Revisits={evidence.get('revisit_count', '?')}]"
            )
        if behavior == "POSSIBLE_FALL":
            alert_msg = f"[{ts}] ⚠ {class_name.upper()} #{track_id} — POSSIBLE FALL DETECTED ({severity}) [Posture=LYING_DOWN]\n"
        elif behavior in ("LOITERING", "SUSPICIOUS_LOITERING"):
            alert_msg = f"[{ts}] ⚠ {class_name.upper()} #{track_id} — SUSPICIOUS LOITERING ({severity}){evidence_note}\n"
        elif behavior == "REPEATED_MOVEMENT":
            alert_msg = f"[{ts}] ⚠ {class_name.upper()} #{track_id} — REPEATED MOVEMENT ({severity}){evidence_note}\n"
        elif behavior == "LONG_DWELL":
            alert_msg = f"[{ts}] ⚠ {class_name.upper()} #{track_id} — LONG DWELL / OBSERVING ({severity}){evidence_note}\n"
        else:
            alert_msg = f"[{ts}] ⚠ {class_name.upper()} #{track_id} — {event_name} ({severity}){evidence_note}\n"

        # Log into primary alert log (bottom text box)
        alert_log = getattr(self, f"{stream_type}_alert_log", None)
        if alert_log:
            self._write_log(alert_log, alert_msg)

        # Log into tactical HUD events log (sidebar)
        events_log = getattr(self, f"{stream_type}_events_log", None)
        if events_log:
            self._write_log(events_log, f"[{ts}] {class_name.upper()} #{track_id} {behavior}\n")

        # Update stats
        self._update_alert_stats(behavior.lower(), f"{class_name.upper()}#{track_id}", severity)
        self._record_incident(stream_type, alert, alert_msg)

    def _write_log(self, log_widget, msg):
        log_widget.config(state="normal")
        log_widget.insert("end", msg)
        log_widget.see("end")
        log_widget.config(state="disabled")

    def _play_alert_sound(self):
        """Start the shared alert sound loop if not already running."""
        if self._alert_sound_active:
            return  # already looping
        self._alert_sound_active = True

        def sound_loop():
            while self._alert_sound_active:
                try:
                    if HAS_WINSOUND:
                        winsound.Beep(2000, 150)
                        time.sleep(0.04)
                        winsound.Beep(2500, 150)
                        time.sleep(0.04)
                        winsound.Beep(3000, 250)
                    else:
                        print('\a', end='', flush=True)
                except Exception:
                    pass
                for _ in range(30):
                    if not self._alert_sound_active:
                        return
                    time.sleep(0.1)

        self._alert_sound_thread = threading.Thread(target=sound_loop, daemon=True)
        self._alert_sound_thread.start()

    def _stop_alert_sound(self):
        """Stop the shared alert sound loop."""
        self._alert_sound_active = False
        self._alert_sound_thread = None

    def _save_incident_snapshot(self, frame, behavior, track_id, severity):
        if frame is None or not hasattr(frame, "shape") or frame.size == 0:
            return None
        incidents_dir = os.path.join(BASE_DIR, "incidents")
        os.makedirs(incidents_dir, exist_ok=True)
        ts = time.strftime("%Y%m%d_%H%M%S")
        filename = f"incident_{behavior}_{track_id}_{severity}_{ts}.jpg"
        filepath = os.path.join(incidents_dir, filename)
        try:
            snapshot = frame.copy()
            label = f"{behavior} | {severity} | ID: {track_id} | {time.strftime('%Y-%m-%d %H:%M:%S')}"
            cv2.putText(snapshot, label, (10, snapshot.shape[0] - 15),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 0, 255), 2)
            cv2.putText(snapshot, "SENTINEL AI — INCIDENT RECORD", (10, 25),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.55, (0, 0, 255), 2)
            success = cv2.imwrite(filepath, snapshot)
            if not success:
                print(f"[WARN] cv2.imwrite failed for {filepath}")
                return None
            return filepath
        except Exception as e:
            print(f"[ERROR] Snapshot save failed: {e}")
        return None

    def _build_stats_tab(self):
        self.stats_vars = {
            "total_persons": tk.IntVar(value=0),
            "total_loiters": tk.IntVar(value=0),      # Loitering
            "total_running": tk.IntVar(value=0),      # Rapid Movement
            "total_pacing": tk.IntVar(value=0),       # Repeated Movement
            "total_falls": tk.IntVar(value=0),        # Possible Fall Events
            "session_start": time.time(),
        }
        self.alert_history = []  # log of alert events

        header = tk.Frame(self.tab_stats, bg=PANEL, height=50)
        header.pack(fill="x", padx=10, pady=(10, 5))
        header.pack_propagate(False)
        tk.Label(header, text="SESSION STATISTICS", font=FONT_HEADING,
                 fg=ACCENT, bg=PANEL).pack(side="left", padx=15, pady=10)

        btn_frame = tk.Frame(self.tab_stats, bg=BG)
        btn_frame.pack(fill="x", padx=10, pady=5)
        tk.Button(btn_frame, text="⟳ REFRESH", font=FONT_BODY, bg=CARD, fg=ACCENT,
                  relief="flat", cursor="hand2",
                  command=self._refresh_stats).pack(side="left", padx=5)
        tk.Button(btn_frame, text="⬇ EXPORT PDF REPORT", font=FONT_BODY,
                  bg=CARD, fg=WARNING, relief="flat", cursor="hand2",
                  command=self._export_pdf).pack(side="left", padx=5)

        cards = tk.Frame(self.tab_stats, bg=BG)
        cards.pack(fill="x", padx=10, pady=10)
        
        stat_items = [
            ("PERSONS TRACKED", "total_persons", "#44aaff"),
            ("LOITERING ALERTS", "total_loiters", WARNING),
            ("RAPID MOVEMENT", "total_running", WARNING),
            ("REPEATED PACING", "total_pacing", WARNING),
            ("POSSIBLE FALLS", "total_falls", WARNING),
        ]
        
        # Grid layout with 4 columns
        for i, (label, key, color) in enumerate(stat_items):
            row = i // 4
            col = i % 4
            card = tk.Frame(cards, bg=CARD, padx=20, pady=12)
            card.grid(row=row, column=col, sticky="nsew", padx=5, pady=5)
            tk.Label(card, textvariable=self.stats_vars[key], font=("Consolas", 24, "bold"),
                     fg=color, bg=CARD).pack()
            tk.Label(card, text=label, font=FONT_SMALL, fg=MUTED, bg=CARD).pack()
            
        for c in range(4):
            cards.grid_columnconfigure(c, weight=1)

        tk.Label(self.tab_stats, text="ALERT HISTORY", font=FONT_BODY,
                 fg=MUTED, bg=BG).pack(padx=15, anchor="w", pady=(10, 2))
        self.stats_log = tk.Text(self.tab_stats, bg=CARD, fg=TEXT, font=FONT_SMALL,
                                  relief="flat", padx=10, pady=5, highlightthickness=0)
        self.stats_log.pack(padx=10, pady=(0, 10), fill="both", expand=True)
        self.stats_log.config(state="disabled")

    def _refresh_stats(self):
        vid_ents = getattr(self, "vid_tracked_entities", {})
        cam_ents = getattr(self, "cam_tracked_entities", {})
        
        persons = sum(1 for e in vid_ents.values() if e["class_name"] == "person") + \
                  sum(1 for e in cam_ents.values() if e["class_name"] == "person")
                  
        self.stats_vars["total_persons"].set(persons)

    def _update_alert_stats(self, alert_type, alert_id, severity=""):
        # Increment behavior counts
        if "loiter" in alert_type:
            self.stats_vars["total_loiters"].set(self.stats_vars["total_loiters"].get() + 1)
        elif "running" in alert_type:
            self.stats_vars["total_running"].set(self.stats_vars["total_running"].get() + 1)
        elif "pacing" in alert_type or "movement" in alert_type:
            self.stats_vars["total_pacing"].set(self.stats_vars["total_pacing"].get() + 1)
        elif "fall" in alert_type:
            self.stats_vars["total_falls"].set(self.stats_vars["total_falls"].get() + 1)

        ts = time.strftime("%H:%M:%S")
        entry = f"[{ts}] {alert_type.upper()} on {alert_id} | SEVERITY: {severity}\n"
        self.alert_history.append(entry)
        if hasattr(self, 'stats_log'):
            self.stats_log.config(state="normal")
            self.stats_log.insert("end", entry)
            self.stats_log.see("end")
            self.stats_log.config(state="disabled")

    def _export_pdf(self):
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib import colors
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.styles import getSampleStyleSheet

            save_path = filedialog.asksaveasfilename(
                defaultextension=".pdf",
                filetypes=[("PDF files", "*.pdf")],
                initialfile=f"sentinel_report_{time.strftime('%Y%m%d_%H%M%S')}.pdf"
            )
            if not save_path:
                return

            doc = SimpleDocTemplate(save_path, pagesize=A4)
            styles = getSampleStyleSheet()
            story = []

            story.append(Paragraph("SENTINEL AI — Threat Detection Session Report", styles["Title"]))
            story.append(Paragraph(f"Generated: {time.strftime('%Y-%m-%d %H:%M:%S')}", styles["Normal"]))
            story.append(Spacer(1, 20))

            duration = int(time.time() - self.stats_vars["session_start"])
            summary_data = [
                ["Metric / Behavior", "Value"],
                ["Session Duration", f"{duration // 60}m {duration % 60}s"],
                ["Persons Tracked", str(self.stats_vars["total_persons"].get())],
                ["Loitering Incidents", str(self.stats_vars["total_loiters"].get())],
                ["Rapid Movement Detections", str(self.stats_vars["total_running"].get())],
                ["Repeated Pacing Incidents", str(self.stats_vars["total_pacing"].get())],
                ["Possible Falls Detected", str(self.stats_vars["total_falls"].get())],
            ]
            t = Table(summary_data, colWidths=[250, 200])
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.black),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
            ]))
            story.append(t)
            story.append(Spacer(1, 20))

            story.append(Paragraph("Suspicious Behavior Alert Log", styles["Heading2"]))
            for entry in self.alert_history:
                story.append(Paragraph(entry.strip(), styles["Normal"]))

            snapshots = [
                incident.get("snapshot_path")
                for incident in self.incidents
                if incident.get("snapshot_path")
            ]
            if snapshots:
                story.append(Spacer(1, 20))
                story.append(Paragraph(f"Session Incident Snapshots: {len(snapshots)} saved for this run", styles["Normal"]))

            doc.build(story)
            self._set_status(f"REPORT EXPORTED: {save_path}", ACCENT)
            if hasattr(self, "stats_log"):
                self._write_log(self.stats_log, f"Report exported: {save_path}\n")
        except Exception as e:
            self._set_status(f"EXPORT FAILED: {e}", DANGER)
            if hasattr(self, "stats_log"):
                self._write_log(self.stats_log, f"Export failed: {e}\n")


# ═══════════════════════════════════════════════════════════════════════════════
# ENTRY POINT
# ═══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    root = tk.Tk()
    app = SentinelApp(root)
    root.mainloop()
